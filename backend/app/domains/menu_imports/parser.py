from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.domains.menu_imports.exceptions import MenuImportValidationError

PARSER_VERSION = "menu-xlsx-v1"
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_WORKSHEETS = 10
MAX_ROWS = 500
MAX_COLUMNS = 50
MAX_HEADER_SCAN_ROWS = 30


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip().casefold()


def display_text(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def safe_filename(value: str) -> str:
    return Path(value.replace("\\", "/")).name


@dataclass(frozen=True)
class ParsedMenuLine:
    line_key: str
    menu_date: date
    meal_name: str
    meal_sort_order: int
    column_name: str
    column_sort_order: int
    original_import_name: str
    normalized_name: str
    source_row: int
    source_column: int
    diner_count: int = 1


@dataclass(frozen=True)
class ParsedMenuLayoutRow:
    meal_name: str
    meal_sort_order: int
    column_name: str
    column_sort_order: int


@dataclass(frozen=True)
class ParsedMenuWorkbook:
    source_hash: str
    parser_version: str
    sheet_name: str
    start_date: date
    end_date: date
    dates: tuple[date, ...]
    meal_names: tuple[str, ...]
    column_count: int
    layout: tuple[ParsedMenuLayoutRow, ...]
    lines: tuple[ParsedMenuLine, ...]
    fatal_errors: tuple[str, ...]
    warnings: tuple[str, ...]


def _validate_archive(payload: bytes) -> None:
    if not is_zipfile(BytesIO(payload)):
        raise MenuImportValidationError("無法讀取 XLSX 檔案")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_SIZE:
                raise MenuImportValidationError("XLSX 解壓縮內容超過安全限制")
    except (BadZipFile, OSError) as exc:
        raise MenuImportValidationError("無法讀取 XLSX 檔案") from exc


def _header(sheet):
    for row in range(1, min(sheet.max_row, MAX_HEADER_SCAN_ROWS) + 1):
        values = [normalize_text(sheet.cell(row, column).value) if sheet.cell(row, column).value is not None else ""
                  for column in range(1, min(sheet.max_column, MAX_COLUMNS) + 1)]
        meal_columns = [index + 1 for index, value in enumerate(values) if value == "餐別"]
        menu_columns = [index + 1 for index, value in enumerate(values) if value == "菜單欄位"]
        if meal_columns and menu_columns:
            meal_column, menu_column = meal_columns[0], menu_columns[0]
            if menu_column != meal_column + 1:
                raise MenuImportValidationError("餐別與菜單欄位必須是相鄰欄位")
            return row, meal_column, menu_column
    raise MenuImportValidationError("找不到必要欄位：餐別、菜單欄位")


def _date_value(cell) -> date:
    if cell.data_type == "f":
        raise MenuImportValidationError("日期標題不可使用公式")
    value = cell.value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = unicodedata.normalize("NFKC", value).strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    raise MenuImportValidationError("七個日期欄必須使用有效的完整日期")


def _dates(sheet, header_row: int, menu_column: int) -> tuple[date, ...]:
    date_cells = [sheet.cell(header_row, menu_column + offset) for offset in range(1, 8)]
    dates = tuple(_date_value(cell) for cell in date_cells)
    extras = [sheet.cell(header_row, column).value for column in range(menu_column + 8, sheet.max_column + 1)]
    if any(value is not None and display_text(value) for value in extras):
        raise MenuImportValidationError("日期欄必須剛好為 7 天")
    if len(set(dates)) != 7:
        raise MenuImportValidationError("日期不可重複")
    if any(dates[index] >= dates[index + 1] for index in range(6)):
        raise MenuImportValidationError("日期必須由左至右遞增")
    if any(dates[index] + timedelta(days=1) != dates[index + 1] for index in range(6)):
        raise MenuImportValidationError("日期必須是連續 7 天")
    return dates


def _merge_map(sheet) -> dict[tuple[int, int], object]:
    result = {}
    for merged_range in sheet.merged_cells.ranges:
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                result[(row, column)] = merged_range
    return result


def _structure_value(sheet, row: int, column: int, merges: dict, label: str):
    merged_range = merges.get((row, column))
    if merged_range is None:
        cell = sheet.cell(row, column)
        if cell.data_type == "f":
            raise MenuImportValidationError(f"{label}不可使用公式")
        return cell.value, None, 0
    if merged_range.min_col != column or merged_range.max_col != column:
        raise MenuImportValidationError(f"{label}只允許向下合併儲存格")
    top = sheet.cell(merged_range.min_row, column)
    if top.data_type == "f":
        raise MenuImportValidationError(f"{label}不可使用公式")
    return top.value, merged_range, row - merged_range.min_row


def _line_key(sheet_name: str, line_date: date, meal_name: str, column_name: str,
              row: int, column: int) -> str:
    payload = json.dumps({
        "parser": PARSER_VERSION, "sheet": sheet_name, "date": line_date.isoformat(),
        "meal": normalize_text(meal_name), "column": normalize_text(column_name),
        "row": row, "source_column": column,
    }, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def parse_menu_xlsx(payload: bytes, filename: str, sheet_name: str | None = None) -> ParsedMenuWorkbook:
    safe_name = safe_filename(filename)
    if Path(safe_name).suffix.lower() != ".xlsx":
        raise MenuImportValidationError("只允許 .xlsx 檔案")
    if not payload:
        raise MenuImportValidationError("XLSX 檔案不可為空")
    if len(payload) > MAX_FILE_SIZE:
        raise MenuImportValidationError("XLSX 檔案超過 10 MB")
    _validate_archive(payload)
    try:
        workbook = load_workbook(BytesIO(payload), data_only=False, keep_vba=False, keep_links=False)
    except Exception as exc:
        raise MenuImportValidationError("無法讀取 XLSX 檔案") from exc
    if not workbook.sheetnames:
        raise MenuImportValidationError("XLSX 沒有工作表")
    if len(workbook.sheetnames) > MAX_WORKSHEETS:
        raise MenuImportValidationError(f"XLSX 工作表不可超過 {MAX_WORKSHEETS} 張")
    if sheet_name is not None:
        if sheet_name not in workbook.sheetnames:
            raise MenuImportValidationError("指定的工作表不存在")
        sheet = workbook[sheet_name]
        header_row, meal_column, menu_column = _header(sheet)
    elif len(workbook.worksheets) == 1:
        sheet = workbook.active
        header_row, meal_column, menu_column = _header(sheet)
    else:
        candidates = []
        for candidate in workbook.worksheets:
            try:
                candidates.append((candidate, _header(candidate)))
            except MenuImportValidationError:
                continue
        if not candidates:
            raise MenuImportValidationError("找不到必要欄位：餐別、菜單欄位")
        if len(candidates) > 1:
            raise MenuImportValidationError("多個工作表符合格式，請指定 sheet_name")
        sheet, (header_row, meal_column, menu_column) = candidates[0]
    if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
        raise MenuImportValidationError(f"工作表不可超過 {MAX_ROWS} 列、{MAX_COLUMNS} 欄")
    dates = _dates(sheet, header_row, menu_column)
    merges = _merge_map(sheet)
    fatal_errors: list[str] = []
    warnings: list[str] = []
    meal_order: dict[str, int] = {}
    meal_display: dict[str, str] = {}
    columns: dict[str, dict[str, int]] = {}
    column_display: dict[tuple[str, str], str] = {}
    column_rows: dict[tuple[str, str], int] = {}
    parsed_lines: list[ParsedMenuLine] = []
    meals_with_dishes: set[str] = set()
    occupied: set[tuple[date, str, str]] = set()
    warned_dish_merge = False
    for row in range(header_row + 1, sheet.max_row + 1):
        dish_cells = [sheet.cell(row, menu_column + offset) for offset in range(1, 8)]
        raw_meal, meal_merge, _ = _structure_value(sheet, row, meal_column, merges, "餐別")
        raw_column, column_merge, column_offset = _structure_value(sheet, row, menu_column, merges, "菜單欄位")
        has_structure = raw_meal is not None or raw_column is not None
        has_dish = any(cell.value is not None and not isinstance(cell, MergedCell) and display_text(cell.value)
                       for cell in dish_cells)
        if not has_structure and not has_dish:
            continue
        if raw_meal is None or not display_text(raw_meal):
            fatal_errors.append(f"第 {row} 列缺少餐別")
            continue
        if raw_column is None or not display_text(raw_column):
            fatal_errors.append(f"第 {row} 列缺少菜單欄位")
            continue
        meal_name = display_text(raw_meal)
        meal_key = normalize_text(meal_name)
        if meal_key not in meal_order:
            meal_order[meal_key] = len(meal_order) + 1
            meal_display[meal_key] = meal_name
            columns[meal_key] = {}
        base_column_name = display_text(raw_column)
        column_name = (f"{base_column_name}{column_offset + 1}"
                       if column_merge is not None and column_merge.max_row > column_merge.min_row
                       else base_column_name)
        column_key = normalize_text(column_name)
        row_key = (meal_key, column_key)
        previous_row = column_rows.get(row_key)
        if previous_row is not None and previous_row != row:
            fatal_errors.append(f"餐別「{meal_name}」展開後欄位「{column_name}」名稱衝突")
        else:
            column_rows[row_key] = row
        if column_key not in columns[meal_key]:
            columns[meal_key][column_key] = len(columns[meal_key]) + 1
            column_display[(meal_key, column_key)] = column_name
        for offset, (line_date, cell) in enumerate(zip(dates, dish_cells, strict=True), 1):
            merged_range = merges.get((row, menu_column + offset))
            if isinstance(cell, MergedCell):
                if not warned_dish_merge:
                    warnings.append("菜色區合併儲存格只讀取左上角，不向其他儲存格展開")
                    warned_dish_merge = True
                continue
            if merged_range is not None and not warned_dish_merge:
                warnings.append("菜色區合併儲存格只讀取左上角，不向其他儲存格展開")
                warned_dish_merge = True
            if cell.data_type == "f":
                fatal_errors.append(f"第 {row} 列第 {menu_column + offset} 欄菜色不可使用公式")
                continue
            if cell.value is None or not display_text(cell.value):
                continue
            original_name = str(cell.value)
            normalized_name = normalize_text(original_name)
            slot_key = (line_date, meal_key, column_key)
            if slot_key in occupied:
                fatal_errors.append(
                    f"{line_date.isoformat()}／{meal_name}／{column_name} 出現兩道菜，無法放入同一菜單欄位"
                )
                continue
            occupied.add(slot_key)
            meals_with_dishes.add(meal_key)
            parsed_lines.append(ParsedMenuLine(
                line_key=_line_key(sheet.title, line_date, meal_name, column_name, row, menu_column + offset),
                menu_date=line_date, meal_name=meal_display[meal_key], meal_sort_order=meal_order[meal_key],
                column_name=column_display[(meal_key, column_key)],
                column_sort_order=columns[meal_key][column_key], original_import_name=original_name,
                normalized_name=normalized_name, source_row=row, source_column=menu_column + offset,
            ))
    for meal_key, meal_name in meal_display.items():
        if meal_key not in meals_with_dishes:
            fatal_errors.append(f"餐別「{meal_name}」沒有任何菜色")
    if not meal_order:
        fatal_errors.append("找不到任何餐別")
    if meal_order and not any(columns.values()):
        fatal_errors.append("找不到任何菜單欄位")
    lines = tuple(sorted(parsed_lines, key=lambda item: (
        item.menu_date, item.meal_sort_order, item.column_sort_order, item.source_row, item.source_column,
    )))
    layout = tuple(
        ParsedMenuLayoutRow(
            meal_name=meal_display[meal_key], meal_sort_order=meal_order[meal_key],
            column_name=column_display[(meal_key, column_key)], column_sort_order=column_order,
        )
        for meal_key in sorted(meal_order, key=meal_order.get)
        for column_key, column_order in sorted(columns[meal_key].items(), key=lambda item: item[1])
    )
    return ParsedMenuWorkbook(
        source_hash=hashlib.sha256(payload).hexdigest(), parser_version=PARSER_VERSION,
        sheet_name=sheet.title, start_date=dates[0], end_date=dates[-1], dates=dates,
        meal_names=tuple(meal_display[key] for key in sorted(meal_order, key=meal_order.get)),
        column_count=sum(len(values) for values in columns.values()), layout=layout, lines=lines,
        fatal_errors=tuple(dict.fromkeys(fatal_errors)), warnings=tuple(dict.fromkeys(warnings)),
    )
