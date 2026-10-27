import re
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.domains.exports.exceptions import SupplierWeeklyExportError
from app.domains.exports.safety import safe_cell_text
from app.shared.domain.quantities import classify_erp_unit, convert_quantity, normalize_unit, quantize_quantity


INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")
THIN_BLACK = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
BLANK_ENTRY_ROWS = 5
TRADITIONAL_CHINESE_WEEKDAYS = (
    "星期一",
    "星期二",
    "星期三",
    "星期四",
    "星期五",
    "星期六",
    "星期日",
)


def supplier_weekly_sheet_title(menu_name: str, requirement_date: date, used: set[str]) -> str:
    """Build an Excel-safe, case-insensitively unique title while retaining MMDD."""
    clean_name = INVALID_SHEET_CHARACTERS.sub("_", menu_name).strip(" '") or "菜單"
    date_suffix = f"_{requirement_date:%m%d}"
    sequence = 1
    while True:
        unique_suffix = "" if sequence == 1 else f"_{sequence}"
        available = 31 - len(date_suffix) - len(unique_suffix)
        candidate = f"{clean_name[:available]}{date_suffix}{unique_suffix}"
        key = candidate.casefold()
        if key not in used:
            used.add(key)
            return candidate
        sequence += 1


def _excel_number(value: Decimal):
    return int(value) if value == value.to_integral_value() else float(value)


def _delivery_date_label(value: date) -> str:
    weekday = TRADITIONAL_CHINESE_WEEKDAYS[value.weekday()]
    return f"送貨日期：{value.year} 年 {value.month} 月 {value.day} 日（{weekday}）"


def _source_value(source: dict, key: str):
    value = source.get(key)
    if value in (None, ""):
        raise SupplierWeeklyExportError("SNAPSHOT_SOURCE_DETAIL_INCOMPLETE", field=key)
    return value


def _historical_meal_name(value, *, required=False) -> str:
    if required and (not isinstance(value, str) or not value.strip()):
        raise SupplierWeeklyExportError("SUPPLIER_ORDER_MEAL_UNAVAILABLE")
    return str(value or "")


def supplier_delivery_date(row: dict, *, breakfast_previous_day=False) -> date:
    requirement_date = row["requirement_date"]
    if breakfast_previous_day:
        meal_name = _historical_meal_name(row.get("meal_type_name"), required=True)
        if "早餐" in meal_name:
            return requirement_date - timedelta(days=1)
    return requirement_date


def rows_from_snapshot_items(items: list[dict], supplier_id, *, require_meal_name=False) -> list[dict]:
    """Expand immutable source rows; never allocate an aggregate adjustment by guessing."""
    result = []
    supplier_key = str(supplier_id)
    matched = False
    for item in items:
        if str(item.get("supplier_id") or "") != supplier_key:
            continue
        matched = True
        sources = item.get("source_summary") or []
        if not sources:
            raise SupplierWeeklyExportError("SNAPSHOT_SOURCE_DETAIL_UNAVAILABLE", item_id=str(item.get("id") or ""))
        target_unit = item.get("purchase_unit_snapshot") or item.get("requirement_unit")
        adjusted_quantity = item.get("adjusted_quantity")
        if target_unit in (None, "") or adjusted_quantity is None:
            raise SupplierWeeklyExportError("SNAPSHOT_EFFECTIVE_QUANTITY_UNAVAILABLE", item_id=str(item.get("id") or ""))
        converted_total = Decimal("0")
        expanded = []
        for source in sources:
            quantity = Decimal(str(_source_value(source, "quantity")))
            unit = str(_source_value(source, "unit"))
            converted = convert_quantity(quantity, unit, target_unit)
            if not converted.convertible or converted.quantity is None:
                raise SupplierWeeklyExportError(
                    "SNAPSHOT_SOURCE_UNIT_CONFLICT",
                    item_id=str(item.get("id") or ""),
                    source_unit=unit,
                    target_unit=target_unit,
                )
            converted_total += converted.quantity
            expanded.append({
                "menu_id": str(_source_value(source, "menu_id")),
                "menu_name": str(_source_value(source, "menu_name")),
                "requirement_date": date.fromisoformat(str(_source_value(source, "requirement_date"))),
                "meal_type_name": _historical_meal_name(source.get("meal_type_name"), required=require_meal_name),
                "dish_id": str(_source_value(source, "dish_id")),
                "dish_name": str(_source_value(source, "dish_name")),
                "ingredient_name": str(item.get("ingredient_name_snapshot") or ""),
                "quantity": converted.quantity,
                "unit": str(target_unit),
                "supplier_name": str(item.get("supplier_name_snapshot") or ""),
            })
        if quantize_quantity(converted_total) != quantize_quantity(Decimal(str(adjusted_quantity))):
            raise SupplierWeeklyExportError(
                "SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE",
                item_id=str(item.get("id") or ""),
            )
        result.extend(expanded)
    if not matched:
        raise SupplierWeeklyExportError("SUPPLIER_ORDER_NO_ITEMS", supplier_id=supplier_key)
    return result


def rows_from_adjustment_lines(lines, supplier_id) -> list[dict]:
    supplier_key = str(supplier_id)
    result = []
    for line in lines:
        if str(line.source_supplier_id or "") != supplier_key:
            continue
        result.append({
            "menu_id": str(line.source_menu_id),
            "menu_name": line.menu_name_snapshot,
            "requirement_date": line.requirement_date,
            "meal_type_name": line.meal_type_name_snapshot,
            "dish_id": str(line.source_dish_id),
            "dish_name": line.dish_name_snapshot,
            "ingredient_name": line.ingredient_name_snapshot,
            "quantity": line.adjusted_quantity if line.adjusted_quantity is not None else line.system_quantity,
            "unit": line.adjusted_unit if line.adjusted_unit is not None else line.system_unit,
            "supplier_name": line.supplier_name_snapshot or "",
            "_sort": (
                line.meal_type_sort_order_snapshot,
                line.menu_dish_sort_order_snapshot,
                line.dish_ingredient_sort_order_snapshot,
                str(line.id),
            ),
        })
    if not result:
        raise SupplierWeeklyExportError("SUPPLIER_ORDER_NO_ITEMS", supplier_id=supplier_key)
    return result


def _export_row(row: dict) -> dict:
    classification = classify_erp_unit(row.get("unit"))
    if classification == "invalid":
        raise SupplierWeeklyExportError(
            "SUPPLIER_ORDER_UNIT_INVALID",
            ingredient_name=str(row.get("ingredient_name") or ""),
            unit=str(row.get("unit") or ""),
        )
    quantity = Decimal(str(row["quantity"]))
    if classification == "weight":
        converted = convert_quantity(quantity, str(row["unit"]), "kg")
        if not converted.convertible or converted.quantity is None:
            raise SupplierWeeklyExportError(
                "SUPPLIER_ORDER_UNIT_INVALID",
                ingredient_name=str(row.get("ingredient_name") or ""),
                unit=str(row.get("unit") or ""),
            )
        return {**row, "_export_quantity": converted.quantity, "_export_unit": "kg"}
    return {**row, "_export_quantity": quantity, "_export_unit": normalize_unit(str(row["unit"]))}


def supplier_weekly_workbook(rows: list[dict], *, breakfast_previous_day=False) -> bytes:
    if not rows:
        raise SupplierWeeklyExportError("SUPPLIER_ORDER_NO_ITEMS")
    rows = [{**_export_row(row), "delivery_date": supplier_delivery_date(row, breakfast_previous_day=breakfast_previous_day)} for row in rows]
    workbook = Workbook()
    workbook.remove(workbook.active)
    grouped = defaultdict(list)
    for row in rows:
        # menu_id is deliberately part of the key: equal display names are distinct menus.
        grouped[(row["delivery_date"], row["menu_id"])].append(row)
    used_titles: set[str] = set()
    for (delivery_date, menu_id), group in sorted(grouped.items(), key=lambda value: (value[0][0], value[1][0]["menu_name"], value[0][1])):
        del menu_id
        first = group[0]
        title = supplier_weekly_sheet_title(first["menu_name"], delivery_date, used_titles)
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:D1")
        sheet["A1"] = "供應商叫貨單"
        sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="000000")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 24
        metadata = [
            f'供應商名稱：{first["supplier_name"] or "未指定供應商"}',
            _delivery_date_label(delivery_date),
            f'菜單名稱：{first["menu_name"]}',
        ]
        for row_index, value in enumerate(metadata, 2):
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=4)
            cell = sheet.cell(row_index, 1, safe_cell_text(value))
            cell.font = Font(name="Arial", size=11, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="center")
        header_row = 5
        headers = ["項次", "訂購品名", "數量", "單位"]
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, value)
            cell.font = Font(name="Arial", size=11, bold=True, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = THIN_BORDER
        ordered = sorted(group, key=lambda row: row.get("_sort") or (row["meal_type_name"], row["dish_name"], row["dish_id"], row["ingredient_name"]))
        for index, row in enumerate(ordered, 1):
            values = [index, safe_cell_text(row["ingredient_name"]), _excel_number(row["_export_quantity"]), row["_export_unit"]]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(header_row + index, column, value)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.font = Font(name="Arial", size=11, color="000000")
                cell.border = THIN_BORDER
                if column == 3:
                    cell.number_format = "General"
        for blank_index in range(1, BLANK_ENTRY_ROWS + 1):
            row_index = header_row + len(ordered) + blank_index
            for column in range(1, 5):
                cell = sheet.cell(row_index, column)
                cell.border = THIN_BORDER
                cell.font = Font(name="Arial", size=11, color="000000")
                cell.alignment = Alignment(horizontal="left", vertical="center")
        for column, width in enumerate((9, 30, 15, 12), 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:5"
        sheet.print_area = f"A1:D{sheet.max_row}"
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
