from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT_NAME = "Microsoft JhengHei"
HEADER_FILL = "D9EAD3"
WARNING_FILL = "FFF2CC"
BORDER_COLOR = "D9D9D9"
WEEKDAY_LABELS = ("星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日")


def _set_run_font(run, size=10, bold=False, color=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT_NAME)


def _set_cell_text(cell, text, *, bold=False, align=None, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    _set_run_font(paragraph.add_run(str(text or "—")), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.insert_element_before(
            shading, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
            "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
            "w:cellMerge", "w:tcPrChange",
        )
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, left=90, bottom=80, right=90):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.insert_element_before(
            margins, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
            "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
        )
    # Word 2007 expects the ECMA-376 transitional left/right margin names.
    # The newer direction-aware start/end elements make Word 2007 reject
    # document.xml even though current Word versions tolerate them.
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_no_wrap(cell):
    properties = cell._tc.get_or_add_tcPr()
    if properties.find(qn("w:noWrap")) is None:
        no_wrap = OxmlElement("w:noWrap")
        properties.insert_element_before(
            no_wrap, "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
            "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
            "w:tcPrChange",
        )


def _set_header_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.insert_element_before(
            borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
            "w:tblPrChange",
        )
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single" if edge == "bottom" else "nil")
        if edge == "bottom":
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), BORDER_COLOR)


def _add_title_row(document, target_date):
    parsed_date = target_date if isinstance(target_date, date) else date.fromisoformat(str(target_date))
    date_label = parsed_date.strftime("%Y/%m/%d")
    weekday_label = WEEKDAY_LABELS[parsed_date.weekday()]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (7.2, 10.4)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, top=25, left=20, bottom=80, right=20)
        _set_no_wrap(cell)
    left, right = table.rows[0].cells
    left.text = ""
    left_paragraph = left.paragraphs[0]
    left_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_paragraph.paragraph_format.space_after = Pt(0)
    _set_run_font(left_paragraph.add_run("月子餐每日異動單"), size=20, bold=True)
    right.text = ""
    right_paragraph = right.paragraphs[0]
    right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_paragraph.paragraph_format.space_after = Pt(0)
    _set_run_font(
        right_paragraph.add_run(f"日期：{date_label}（{weekday_label}）"), size=20, bold=True,
    )
    _set_header_table_borders(table)
    return table


def _style_table(table, widths, *, warning_rows=()):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        _prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            if index < len(widths):
                cell.width = Cm(widths[index])
            _set_cell_margins(cell)
            if row_index == 0:
                _shade_cell(cell, HEADER_FILL)
            elif row_index in warning_rows:
                _shade_cell(cell, WARNING_FILL)
    _repeat_header(table.rows[0])


def _add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_run_font(paragraph.add_run(text), size=13 if level == 1 else 11, bold=True)
    return paragraph


def _dish_text(dish):
    if not dish:
        return "—"
    name = str(dish.get("name") or "").strip()
    return name or "—"


def _case_text(item):
    name = str(item.get("case_name") or "").strip()
    return name or "—"


def _restriction_text(item):
    names = [str(group.get("name") or "").strip() for group in item.get("restriction_groups", [])]
    return "、".join(name for name in names if name) or "—"


def _warning_messages(items):
    return list(dict.fromkeys(str(item.get("message") or "").strip() for item in items if item.get("message")))


def _add_warning_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    _set_run_font(paragraph.add_run(f"⚠ {text}"), size=10, bold=True, color="9C6500")


def _add_replacement_details(document, groups):
    items = [item for group in groups for item in group.get("items", [])]
    if not items:
        return
    _add_heading(document, "個案異動明細", 2)
    table = document.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ("床號", "個案", "原菜", "替代菜", "禁忌")):
        _set_cell_text(cell, value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    warning_rows = []
    for item in items:
        row = table.add_row()
        stale = item.get("status") == "requires_reconfirmation"
        values = (
            item.get("current_room"), _case_text(item), _dish_text(item.get("original_dish")),
            _dish_text(item.get("replacement_dish")), _restriction_text(item),
        )
        for cell, value in zip(row.cells, values):
            _set_cell_text(cell, value)
        if stale:
            _set_cell_text(row.cells[4], f"⚠ 待重新確認\n{_restriction_text(item)}")
            warning_rows.append(len(table.rows) - 1)
    _style_table(table, (1.5, 3.0, 3.7, 3.7, 4.5), warning_rows=warning_rows)


def _add_acknowledgements(document, items):
    if not items:
        return
    _add_heading(document, "人工確認 不需替代", 2)
    table = document.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ("床號", "個案", "原菜", "結果", "備註")):
        _set_cell_text(cell, value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    warning_rows = []
    for item in items:
        stale = item.get("status") == "requires_reconfirmation"
        row = table.add_row()
        values = (
            item.get("current_room"), _case_text(item), _dish_text(item.get("original_dish")),
            "⚠ 待重新確認" if stale else "人工確認，不需替代", item.get("note") or "—",
        )
        for cell, value in zip(row.cells, values):
            _set_cell_text(cell, value)
        if stale:
            warning_rows.append(len(table.rows) - 1)
    _style_table(table, (1.5, 3.0, 4.0, 3.5, 4.4), warning_rows=warning_rows)


def _add_reconfirmation(document, items):
    if not items:
        return
    _add_heading(document, "⚠ 待重新確認", 2)
    _add_warning_paragraph(document, "此替代內容尚需營養師重新確認")
    for item in items:
        treatment = "人工確認不需替代"
        if item.get("handling_type") == "replacement":
            treatment = f"替代為 {_dish_text(item.get('replacement_dish'))}"
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(4)
        lead = f"{item.get('current_room') or '—'}｜{item.get('case_name') or '—'}｜{_dish_text(item.get('original_dish'))}｜原處理：{treatment}"
        _set_run_font(paragraph.add_run(lead), size=9.5, bold=True)
        messages = _warning_messages(item.get("warnings", []))
        if messages:
            _set_run_font(paragraph.add_run("\n" + "；".join(messages)), size=9, color="9C6500")


def render_daily_change_sheet_docx(data: dict) -> bytes:
    """Render the authoritative daily change-sheet read model as editable OOXML."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    target_date = data["target_date"]
    _add_title_row(document, target_date)

    for meal in data["meals"]:
        if not meal.get("has_changes"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
            _set_run_font(paragraph.add_run(f"{meal['meal_label']}　"), size=11, bold=True)
            _set_run_font(paragraph.add_run("無異動"), size=10, color="666666")
            continue
        _add_heading(document, meal["meal_label"], 1)
        _add_replacement_details(document, meal.get("replacement_groups", []))
        _add_acknowledgements(document, meal.get("manual_acknowledgements", []))
        stale_items = meal.get("requires_reconfirmation", [])
        _add_reconfirmation(document, stale_items)
        # Meal/day warning collections also contain coverage diagnostics from the
        # conflict engine. They remain in the authoritative API, but are not
        # operational instructions for the kitchen change sheet. Genuine stale
        # reasons are rendered above from the flattened reconfirmation items.

    output = BytesIO()
    document.save(output)
    return output.getvalue()
