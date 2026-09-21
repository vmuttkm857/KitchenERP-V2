from __future__ import annotations

import math
import unicodedata
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, Side


FONT_NAME = "標楷體"
ROOM_COLOR = "7030A0"
RESTRICTION_COLOR = "FF0000"
NOTE_COLOR = "1F4E79"
RICE_WINE_COLOR = "FF0000"
NO_HERBAL_COLOR = "00B0F0"
THIN_BORDER = Border(*(Side(style="thin", color="000000") for _ in range(4)))
COLUMN_WIDTHS = {
    "A": 27.6328125,
    "B": 28.36328125,
    "C": 58.08984375,
    "D": 29.6328125,
    "E": 31.453125,
    "F": 58.453125,
}


def _roc_year(value) -> int:
    return value.year - 1911


def _font(size: float, *, bold: bool = False, color: str | None = None) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, charset=136, family=4)


def _inline_font(size: float, *, bold: bool = False, color: str | None = None) -> InlineFont:
    return InlineFont(rFont=FONT_NAME, sz=size, b=bold, color=color, charset=136, family=4)


def _preparation_mode(value: object) -> tuple[str, str | None]:
    label = str(value or "").strip()
    if label == "要中藥":
        return "", None
    if label in {"米酒水＋麻油", "米酒水+麻油", "米酒水"}:
        return "米酒水", RICE_WINE_COLOR
    if label == "不中藥":
        return label, NO_HERBAL_COLOR
    return label, None


def _service_start(item: dict) -> str:
    value = item["service_start_date"]
    if hasattr(value, "month"):
        month, day = value.month, value.day
    else:
        _, month, day = str(value).split("-")
        month, day = int(month), int(day)
    return f'{month}/{day}{item["service_start_meal_label"]}起'


def _restriction_text(item: dict) -> str:
    return "、".join(
        group["name"] + ("（已停用）" if not group["is_active"] else "")
        for group in item.get("restriction_groups", [])
    )


def _set_cell_style(cell, *, font: Font, horizontal: str | None = None, wrap: bool = True):
    cell.font = font
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=wrap)
    cell.border = THIN_BORDER


def _restriction_note_value(cell, item: dict):
    restrictions = _restriction_text(item)
    note = str(item.get("service_note") or "").strip()
    if restrictions and note:
        cell.value = CellRichText(
            TextBlock(_inline_font(26, color=RESTRICTION_COLOR), restrictions),
            TextBlock(_inline_font(26, bold=True, color=NOTE_COLOR), f"\n{note}"),
        )
        cell.font = _font(26, color=RESTRICTION_COLOR)
    elif restrictions:
        cell.value = restrictions
        cell.font = _font(26, color=RESTRICTION_COLOR)
    elif note:
        cell.value = note
        cell.font = _font(26, bold=True, color=NOTE_COLOR)
    else:
        cell.value = ""
        cell.font = _font(26, color=RESTRICTION_COLOR)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _display_width(value: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in "WFA" else 1 for char in value)


def _wrapped_lines(value: str, capacity: int = 28) -> int:
    if not value:
        return 0
    return sum(max(1, math.ceil(_display_width(line) / capacity)) for line in value.splitlines())


def _main_row_height(left: dict | None, right: dict | None) -> float:
    line_count = 1
    for item in (left, right):
        if not item:
            continue
        restrictions = _restriction_text(item)
        note = str(item.get("service_note") or "").strip()
        line_count = max(line_count, _wrapped_lines(restrictions), _wrapped_lines(note))
    if line_count <= 3:
        return 252.75
    if line_count <= 5:
        return 292.5
    return 321.75 + max(0, line_count - 7) * 35.5


def _write_case(sheet, main_row: int, start_col: int, item: dict | None):
    room = sheet.cell(main_row, start_col)
    preparation = sheet.cell(main_row, start_col + 1)
    detail = sheet.cell(main_row, start_col + 2)
    name = sheet.cell(main_row + 1, start_col)
    start = sheet.cell(main_row + 1, start_col + 1)
    blank = sheet.cell(main_row + 1, start_col + 2)

    if item:
        room.value = item["current_room"]
        preparation.value, preparation_color = _preparation_mode(item["preparation_mode_label"])
        name.value = item["name"]
        start.value = _service_start(item)
    else:
        room.value = preparation.value = name.value = start.value = ""
        preparation_color = None

    _set_cell_style(room, font=_font(36, bold=True, color=ROOM_COLOR), horizontal="center")
    _set_cell_style(preparation, font=_font(28, color=preparation_color), horizontal="center")
    _restriction_note_value(detail, item or {})
    _set_cell_style(name, font=_font(18), horizontal="center", wrap=False)
    _set_cell_style(start, font=_font(24), horizontal="center", wrap=False)
    _set_cell_style(blank, font=_font(26, bold=True), horizontal="center", wrap=False)


def render_catering_overview_xlsx(data: dict) -> bytes:
    """Render the existing catering-overview read model in the legacy paired layout."""
    target_date = data["target_date"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{_roc_year(target_date)}{target_date:%m%d}"
    sheet.sheet_format.defaultRowHeight = 20.5
    sheet.sheet_view.zoomScale = 25
    sheet.sheet_view.zoomScaleNormal = 25

    for column, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"製表日期{_roc_year(target_date)}/{target_date.month}/{target_date.day}"
    sheet["A1"].font = _font(16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    items = list(data.get("items", []))
    for pair_index in range(0, len(items), 2):
        main_row = 2 + pair_index
        left = items[pair_index]
        right = items[pair_index + 1] if pair_index + 1 < len(items) else None
        _write_case(sheet, main_row, 1, left)
        _write_case(sheet, main_row, 4, right)
        sheet.row_dimensions[main_row].height = _main_row_height(left, right)
        sheet.row_dimensions[main_row + 1].height = 35.5

    sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 40
    sheet.page_margins.left = 0.3937007874015748
    sheet.page_margins.right = 0
    sheet.page_margins.top = 0
    sheet.page_margins.bottom = 0
    sheet.page_margins.header = 0.11811023622047245
    sheet.page_margins.footer = 0.11811023622047245

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
