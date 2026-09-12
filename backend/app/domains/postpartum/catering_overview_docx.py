from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.domains.postpartum.catering_overview import paginate_catering_items


FONT_NAME = "Microsoft JhengHei"


def _font(run, size=12, bold=False, color=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT_NAME)


def _cell(cell, text, *, size=11, bold=False, color=None, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    _font(paragraph.add_run(str(text or "—")), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def _remove_table_borders(table):
    props = table._tbl.tblPr
    borders = props.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _paragraph_bottom_border(paragraph, color="D9D9D9"):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        props.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _row_no_split(row, repeat=False):
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        props.append(header)


def _set_width(cell, width):
    cell.width = Cm(width)
    props = cell._tc.get_or_add_tcPr()
    tc_width = props.find(qn("w:tcW"))
    if tc_width is not None:
        tc_width.set(qn("w:w"), str(int(width * 567)))
        tc_width.set(qn("w:type"), "dxa")


def _no_wrap(cell):
    props = cell._tc.get_or_add_tcPr()
    if props.find(qn("w:noWrap")) is None:
        props.append(OxmlElement("w:noWrap"))


def _page_heading(document, data, page_number, page_count):
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (9.8, 5.5, 4.3)
    value = data["target_date"]
    if hasattr(value, "strftime"):
        date_text = value.strftime("%Y 年 %m 月 %d 日")
    else:
        year, month, day = str(value).split("-")
        date_text = f"{year} 年 {month} 月 {day} 日"
    values = (
        (f"{date_text}（{data['weekday_label']}）", 15, True),
        (f"供餐總床數：{data['total']} 床", 13, True),
        (f"第 {page_number} 頁 / 共 {page_count} 頁", 11, False),
    )
    for cell, width, content in zip(table.rows[0].cells, widths, values):
        _set_width(cell, width)
        _no_wrap(cell)
        _cell(cell, content[0], size=content[1], bold=content[2], align=WD_ALIGN_PARAGRAPH.CENTER)
    _remove_table_borders(table)
    _row_no_split(table.rows[0])
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = .2


def _service_moment(date_value, meal_label):
    if date_value is None or not meal_label:
        return "—"
    if hasattr(date_value, "month"):
        date_text = f"{date_value.month}/{date_value.day}"
    else:
        _, month, day = str(date_value).split("-")
        date_text = f"{int(month)}/{int(day)}"
    return f"{date_text} {meal_label}"


def _person_cell(cell, item):
    cell.text = ""
    values = (
        (item["name"] or "—", 15, True),
        (f"起伙：{_service_moment(item['service_start_date'], item['service_start_meal_label'])}", 9.5, False),
    )
    for index, (text, size, bold) in enumerate(values):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(1 if index == 1 else 0)
        paragraph.paragraph_format.space_after = Pt(1 if index == 0 else 0)
        paragraph.paragraph_format.line_spacing = 1.0
        _font(paragraph.add_run(text), size=size, bold=bold)
        if index == 0:
            _paragraph_bottom_border(paragraph)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _restriction_note_cell(cell, item):
    cell.text = ""
    restrictions = "、".join(
        group["name"] + ("（已停用）" if not group["is_active"] else "")
        for group in item["restriction_groups"]
    )
    note = str(item.get("service_note") or "").strip()
    parts = []
    if restrictions:
        parts.append((restrictions, "A61B1B", True))
    if note:
        parts.append((note, "1F4E79", True))
    if not parts:
        parts.append(("—", None, False))
    for index, (text, color, bold) in enumerate(parts):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        _font(paragraph.add_run(text), size=12, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(document, items):
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (2.0, 4.3, 3.5, 9.8)
    headers = ("床號", "姓名", "調理方式", "飲食禁忌／備註")
    for cell, label, width in zip(table.rows[0].cells, headers, widths):
        _set_width(cell, width)
        _cell(cell, label, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(cell, "D9EAD3")
    _row_no_split(table.rows[0], repeat=True)
    for item in items:
        row = table.add_row()
        values = (
            (item["current_room"], 15, True, None),
            (None, 12, True, None),
            (item["preparation_mode_label"], 12, False, None),
            (None, 12, False, None),
        )
        for index, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
            _set_width(cell, width)
            if index == 1:
                _person_cell(cell, item)
            elif index == 3:
                _restriction_note_cell(cell, item)
            else:
                _cell(cell, value[0], size=value[1], bold=value[2], color=value[3])
        _row_no_split(row)
    return table


def render_catering_overview_docx(data: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(.8)
    section.bottom_margin = Cm(.8)
    section.left_margin = Cm(.7)
    section.right_margin = Cm(.7)
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    pages = paginate_catering_items(data["items"])
    for index, items in enumerate(pages, start=1):
        if index > 1:
            document.add_page_break()
        _page_heading(document, data, index, len(pages))
        if items:
            _table(document, items)
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(28)
            _font(paragraph.add_run("本日沒有符合供餐條件的月子餐個案"), size=14, bold=True, color="68766D")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
