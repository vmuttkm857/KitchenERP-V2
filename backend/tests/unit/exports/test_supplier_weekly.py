from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.api.v1.routes.exports import mapped
from app.domains.exports.exceptions import SupplierWeeklyExportError
from app.domains.exports.supplier_weekly import (
    rows_from_adjustment_lines,
    rows_from_snapshot_items,
    supplier_weekly_sheet_title,
    supplier_weekly_workbook,
)


def source(menu_id, menu_name, day, dish_id, dish_name, quantity="1.000000", unit="kg"):
    return {
        "menu_id": str(menu_id),
        "menu_name": menu_name,
        "requirement_date": day,
        "meal_type_name": "中餐",
        "dish_id": str(dish_id),
        "dish_name": dish_name,
        "quantity": quantity,
        "unit": unit,
    }


def item(supplier_id, sources, adjusted="3.000000"):
    return {
        "id": str(uuid4()),
        "supplier_id": str(supplier_id),
        "supplier_name_snapshot": "方豪",
        "ingredient_name_snapshot": "高麗菜",
        "requirement_unit": "kg",
        "purchase_unit_snapshot": "kg",
        "adjusted_quantity": adjusted,
        "source_summary": sources,
    }


def test_snapshot_sources_keep_same_ingredient_in_different_dishes_as_distinct_rows():
    supplier_id = uuid4()
    menu_id = uuid4()
    rows = rows_from_snapshot_items([
        item(supplier_id, [
            source(menu_id, "治療伙食單-2", "2026-11-02", uuid4(), "炒高麗菜", "1.000000"),
            source(menu_id, "治療伙食單-2", "2026-11-02", uuid4(), "高麗菜湯", "2.000000"),
        ])
    ], supplier_id)
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows)))
    assert workbook.sheetnames == ["治療伙食單-2_1102"]
    sheet = workbook.active
    assert sheet["A1"].value == "供應商叫貨單"
    assert sheet["A2"].value == "供應商名稱：方豪"
    assert sheet["A3"].value == "送貨日期：2026 年 11 月 2 日（星期一）"
    assert sheet["A4"].value == "菜單名稱：治療伙食單-2"
    assert [sheet.cell(row, 2).value for row in (6, 7)] == ["高麗菜", "高麗菜"]
    assert [sheet.cell(row, 3).value for row in (6, 7)] == [1, 2]
    assert [sheet.cell(5, column).value for column in range(1, 5)] == ["項次", "訂購品名", "數量", "單位"]
    assert sheet.auto_filter.ref is None
    assert all(sheet.cell(5, column).fill.fill_type is None for column in range(1, 5))
    assert sheet.page_setup.paperSize == 9 and sheet.page_setup.orientation == "portrait"
    assert sheet.max_row == 12  # two actual rows plus five bordered manual-entry rows
    assert sheet["A1"].alignment.horizontal == "center"
    assert all(sheet.cell(5, column).alignment.horizontal == "left" for column in range(1, 5))
    assert all(sheet.cell(row, column).alignment.horizontal == "left" for row in range(6, 13) for column in range(1, 5))
    all_text = "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "醫院名稱" not in all_text and "未設定" not in all_text
    assert "菜餚說明" not in all_text and "是否裁切" not in all_text and "廚房" not in all_text


def test_workbook_groups_by_date_and_menu_id_not_display_name():
    supplier_id = uuid4()
    rows = rows_from_snapshot_items([
        item(supplier_id, [source(uuid4(), "同名菜單", "2026-11-02", uuid4(), "菜一")], "1.000000"),
        item(supplier_id, [source(uuid4(), "同名菜單", "2026-11-02", uuid4(), "菜二")], "1.000000"),
    ], supplier_id)
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows)))
    assert workbook.sheetnames == ["同名菜單_1102", "同名菜單_1102_2"]
    assert all(sheet["A4"].value == "菜單名稱：同名菜單" for sheet in workbook.worksheets)


def test_each_sheet_uses_its_own_delivery_date_for_traditional_chinese_weekday():
    menu_id = str(uuid4())
    common = {
        "menu_id": menu_id,
        "menu_name": "星期驗證菜單",
        "meal_type_name": "中餐",
        "supplier_name": "方豪",
        "ingredient_name": "高麗菜",
        "quantity": Decimal("1"),
        "unit": "kg",
    }
    rows = [
        {**common, "requirement_date": day, "dish_id": str(uuid4()), "dish_name": f"菜色-{day}"}
        for day in (date(2026, 11, 2), date(2026, 11, 3), date(2026, 11, 8))
    ]
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows)))
    assert [sheet["A3"].value for sheet in workbook.worksheets] == [
        "送貨日期：2026 年 11 月 2 日（星期一）",
        "送貨日期：2026 年 11 月 3 日（星期二）",
        "送貨日期：2026 年 11 月 8 日（星期日）",
    ]


def test_sheet_title_keeps_date_and_uniqueness_when_menu_name_is_long_or_invalid():
    used = set()
    first = supplier_weekly_sheet_title("超長/菜單名稱" * 5, date(2026, 11, 2), used)
    second = supplier_weekly_sheet_title("超長/菜單名稱" * 5, date(2026, 11, 2), used)
    assert len(first) <= 31 and first.endswith("_1102")
    assert len(second) <= 31 and second.endswith("_1102_2")
    assert first != second and "/" not in first


def test_aggregate_adjustment_that_cannot_be_allocated_to_dishes_is_rejected():
    supplier_id = uuid4()
    with pytest.raises(SupplierWeeklyExportError) as caught:
        rows_from_snapshot_items([
            item(supplier_id, [source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜色", "3.000000")], "4.000000")
        ], supplier_id)
    assert caught.value.code == "SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE"


def test_source_units_use_existing_weight_conversion_without_rounding():
    supplier_id = uuid4()
    rows = rows_from_snapshot_items([
        {
            **item(supplier_id, [
                source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜一", "1500", "g"),
                source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜二", "2", "斤"),
            ], "2.700000"),
            "purchase_unit_snapshot": "kg",
        }
    ], supplier_id)
    assert [(row["quantity"], row["unit"]) for row in rows] == [(Decimal("1.5"), "kg"), (Decimal("1.2"), "kg")]


def test_workbook_converts_all_supported_weight_units_to_numeric_kg_without_extra_rounding():
    menu_id = str(uuid4())
    common = {
        "menu_id": menu_id,
        "menu_name": "測試菜單",
        "requirement_date": date(2026, 11, 2),
        "meal_type_name": "中餐",
        "supplier_name": "方豪",
    }
    rows = [
        {**common, "dish_id": str(uuid4()), "dish_name": "菜一", "ingredient_name": "食材一", "quantity": Decimal("4000"), "unit": "g"},
        {**common, "dish_id": str(uuid4()), "dish_name": "菜二", "ingredient_name": "食材二", "quantity": Decimal("6500"), "unit": "g"},
        {**common, "dish_id": str(uuid4()), "dish_name": "菜三", "ingredient_name": "食材三", "quantity": Decimal("5.25"), "unit": "斤"},
        {**common, "dish_id": str(uuid4()), "dish_name": "菜四", "ingredient_name": "食材四", "quantity": Decimal("2.75"), "unit": "kg"},
    ]
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(rows))).active
    exported={sheet.cell(row,2).value:sheet.cell(row,3).value for row in range(6,10)}
    assert exported == {"食材一":4,"食材二":6.5,"食材三":3.15,"食材四":2.75}
    assert [sheet.cell(row, 3).data_type for row in range(6, 10)] == ["n", "n", "n", "n"]
    assert [sheet.cell(row, 4).value for row in range(6, 10)] == ["kg", "kg", "kg", "kg"]


@pytest.mark.parametrize("unit", ["片", "個", "隻", "包", "盒", "袋", "箱", "ml", "L", "罐", "桶"])
def test_workbook_preserves_legal_non_weight_quantity_and_unit(unit):
    common = {
        "menu_id": str(uuid4()),
        "menu_name": "測試菜單",
        "requirement_date": date(2026, 11, 2),
        "meal_type_name": "中餐",
        "supplier_name": "方豪",
    }
    rows = [
        {**common, "dish_id": str(uuid4()), "dish_name": "重量菜色", "ingredient_name": "高麗菜", "quantity": Decimal("4000"), "unit": "g"},
        {**common, "dish_id": str(uuid4()), "dish_name": "非重量菜色", "ingredient_name": f"非重量-{unit}", "quantity": Decimal("10"), "unit": unit},
    ]
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(rows))).active
    actual = {
        sheet.cell(row, 2).value: (sheet.cell(row, 3).value, sheet.cell(row, 4).value, sheet.cell(row, 3).data_type)
        for row in range(6, 8)
    }
    assert actual == {"高麗菜": (4, "kg", "n"), f"非重量-{unit}": (10, unit, "n")}


def test_workbook_with_only_piece_each_and_animal_units_is_exportable():
    menu_id = str(uuid4())
    rows = [
        {
            "menu_id": menu_id,
            "menu_name": "測試菜單",
            "requirement_date": date(2026, 11, 2),
            "meal_type_name": "中餐",
            "dish_id": str(uuid4()),
            "dish_name": "菜色",
            "ingredient_name": ingredient,
            "quantity": quantity,
            "unit": unit,
            "supplier_name": "方豪",
        }
        for ingredient, quantity, unit in (("雞排", Decimal("120"), "片"), ("雞蛋", Decimal("450"), "個"), ("柳葉魚", Decimal("210"), "隻"))
    ]
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(rows))).active
    actual = {(sheet.cell(row, 2).value, sheet.cell(row, 3).value, sheet.cell(row, 4).value) for row in range(6, 9)}
    assert actual == {("雞排", 120, "片"), ("雞蛋", 450, "個"), ("柳葉魚", 210, "隻")}


def test_supplier_weekly_no_items_error_has_a_stable_client_code():
    response = mapped(SupplierWeeklyExportError("SUPPLIER_ORDER_NO_ITEMS"))
    assert response.status_code == 422
    assert response.detail == {"code": "SUPPLIER_ORDER_NO_ITEMS"}


def test_workbook_without_any_demand_rows_returns_clear_no_items_error():
    with pytest.raises(SupplierWeeklyExportError) as caught:
        supplier_weekly_workbook([])
    assert caught.value.code == "SUPPLIER_ORDER_NO_ITEMS"


@pytest.mark.parametrize("unit", [None, "", "   ", "公克", "damaged-unit"])
@pytest.mark.parametrize("breakfast_previous_day", [False, True])
def test_workbook_rejects_missing_or_unknown_units_instead_of_silently_omitting(unit, breakfast_previous_day):
    with pytest.raises(SupplierWeeklyExportError) as caught:
        supplier_weekly_workbook([{
            "menu_id": str(uuid4()),
            "menu_name": "測試菜單",
            "requirement_date": date(2026, 11, 2),
            "meal_type_name": "中餐",
            "dish_id": str(uuid4()),
            "dish_name": "菜色",
            "ingredient_name": "資料異常食材",
            "quantity": Decimal("10"),
            "unit": unit,
            "supplier_name": "方豪",
        }], breakfast_previous_day=breakfast_previous_day)
    assert caught.value.code == "SUPPLIER_ORDER_UNIT_INVALID"


def test_non_weight_only_date_menu_group_creates_its_own_sheet():
    menu_id = str(uuid4())
    rows = [
        {
            "menu_id": menu_id,
            "menu_name": "重量菜單",
            "requirement_date": date(2026, 11, 2),
            "meal_type_name": "中餐",
            "dish_id": str(uuid4()),
            "dish_name": "重量菜色",
            "ingredient_name": "高麗菜",
            "quantity": Decimal("1"),
            "unit": "kg",
            "supplier_name": "方豪",
        },
        {
            "menu_id": menu_id,
            "menu_name": "非重量菜單",
            "requirement_date": date(2026, 11, 3),
            "meal_type_name": "中餐",
            "dish_id": str(uuid4()),
            "dish_name": "非重量菜色",
            "ingredient_name": "雞排",
            "quantity": Decimal("10"),
            "unit": "片",
            "supplier_name": "方豪",
        },
    ]
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows)))
    assert workbook.sheetnames == ["重量菜單_1102", "非重量菜單_1103"]
    assert workbook["非重量菜單_1103"]["C6"].value == 10
    assert workbook["非重量菜單_1103"]["D6"].value == "片"


def test_non_weight_support_does_not_bypass_snapshot_source_reconciliation():
    supplier_id = uuid4()
    bad_item = item(
        supplier_id,
        [source(uuid4(), "菜單", "2026-11-02", uuid4(), "雞排", "10", "片")],
        "11",
    )
    bad_item["requirement_unit"] = "片"
    bad_item["purchase_unit_snapshot"] = "片"
    with pytest.raises(SupplierWeeklyExportError) as caught:
        rows_from_snapshot_items([bad_item], supplier_id)
    assert caught.value.code == "SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE"


def test_adjustment_export_preserves_saved_effective_quantity_and_unit_for_every_legal_unit():
    supplier_id = uuid4()
    common = {
        "source_supplier_id": supplier_id,
        "source_menu_id": uuid4(),
        "menu_name_snapshot": "叫貨調整菜單",
        "requirement_date": date(2026, 11, 2),
        "meal_type_name_snapshot": "中餐",
        "supplier_name_snapshot": "方豪",
        "meal_type_sort_order_snapshot": 1,
        "menu_dish_sort_order_snapshot": 1,
        "dish_ingredient_sort_order_snapshot": 1,
    }
    lines = [
        SimpleNamespace(
            **common,
            id=uuid4(),
            source_dish_id=uuid4(),
            dish_name_snapshot="調整重量菜色",
            ingredient_name_snapshot="高麗菜",
            system_quantity=Decimal("4000"),
            system_unit="g",
            adjusted_quantity=Decimal("5.25"),
            adjusted_unit="斤",
        ),
        SimpleNamespace(
            **common,
            id=uuid4(),
            source_dish_id=uuid4(),
            dish_name_snapshot="調整非重量菜色",
            ingredient_name_snapshot="雞排",
            system_quantity=Decimal("1000"),
            system_unit="g",
            adjusted_quantity=Decimal("10"),
            adjusted_unit="片",
        ),
    ]
    rows = rows_from_adjustment_lines(lines, supplier_id)
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(rows))).active
    actual = {
        sheet.cell(row, 2).value: (sheet.cell(row, 3).value, sheet.cell(row, 4).value)
        for row in range(6, 8)
    }
    assert actual == {"高麗菜": (3.15, "kg"), "雞排": (10, "片")}


def delivery_row(**changes):
    return {
        "menu_id": "menu-a", "menu_name": "測試菜單", "requirement_date": date(2026, 11, 2),
        "meal_type_name": "早餐", "dish_id": "dish-a", "dish_name": "早餐蛋餅",
        "ingredient_name": "雞蛋", "quantity": Decimal("120"), "unit": "個", "supplier_name": "供應商",
        **changes,
    }


def workbook_state(payload):
    workbook = load_workbook(BytesIO(payload))
    return [(
        sheet.title,
        [(cell.coordinate, cell.value, cell.data_type, str(cell.alignment), str(cell.border), str(cell.font), str(cell.fill), cell.number_format)
         for row in sheet for cell in row],
        str(sheet.page_setup), str(sheet.page_margins), str(sheet.print_area), str(sheet.merged_cells),
        [(key, value.width) for key, value in sheet.column_dimensions.items()],
    ) for sheet in workbook]


def test_default_and_explicit_false_preserve_workbook_and_source_even_without_meal():
    rows = [delivery_row(meal_type_name="")]
    original = deepcopy(rows)
    implicit = supplier_weekly_workbook(rows)
    explicit = supplier_weekly_workbook(rows, breakfast_previous_day=False)
    assert workbook_state(implicit) == workbook_state(explicit)
    assert load_workbook(BytesIO(implicit)).sheetnames == ["測試菜單_1102"]
    assert rows == original


@pytest.mark.parametrize("meal_name,expected", [
    ("早餐", "1101"), ("員工早餐", "1101"), ("病患早餐", "1101"), ("早餐（特殊餐）", "1101"),
    ("午餐", "1102"), ("晚餐", "1102"), ("晚點", "1102"), ("早點", "1102"),
])
def test_breakfast_uses_only_historical_meal_name(meal_name, expected):
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(
        [delivery_row(meal_type_name=meal_name)], breakfast_previous_day=True,
    ))).active
    assert sheet.title == f"測試菜單_{expected}"  # dish name contains 早餐 even for non-breakfast meals
    assert sheet["C6"].value == 120 and sheet["D6"].value == "個"


@pytest.mark.parametrize("usage,expected_label,title", [
    (date(2026, 11, 2), "2026 年 11 月 1 日（星期日）", "1101"),
    (date(2026, 12, 1), "2026 年 11 月 30 日（星期一）", "1130"),
    (date(2027, 1, 1), "2026 年 12 月 31 日（星期四）", "1231"),
    (date(2028, 3, 1), "2028 年 2 月 29 日（星期二）", "0229"),
    (date(2027, 3, 1), "2027 年 2 月 28 日（星期日）", "0228"),
])
def test_delivery_dates_cross_week_month_year_and_leap_boundaries(usage, expected_label, title):
    sheet = load_workbook(BytesIO(supplier_weekly_workbook(
        [delivery_row(requirement_date=usage)], breakfast_previous_day=True,
    ))).active
    assert sheet.title == f"測試菜單_{title}"
    assert sheet["A3"].value == f"送貨日期：{expected_label}"


def test_delivery_group_combines_usage_dates_without_merging_dishes_or_menus():
    rows = [
        delivery_row(meal_type_name="午餐", dish_id="lunch", quantity=Decimal("1"), unit="kg"),
        delivery_row(meal_type_name="晚餐", dish_id="dinner", quantity=Decimal("2"), unit="kg"),
        delivery_row(requirement_date=date(2026, 11, 3), quantity=Decimal("3"), unit="kg"),
        delivery_row(menu_id="menu-b", requirement_date=date(2026, 11, 3), quantity=Decimal("4"), unit="kg"),
    ]
    original = deepcopy(rows)
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows, breakfast_previous_day=True)))
    assert workbook.sheetnames == ["測試菜單_1102", "測試菜單_1102_2"]
    actual = [[sheet.cell(i, 3).value for i in range(6, sheet.max_row + 1) if sheet.cell(i, 1).value is not None]
              for sheet in workbook]
    assert sorted(actual[0]) == [1, 2, 3] and actual[1] == [4]
    assert sum(map(sum, actual)) == 10 and sum(map(len, actual)) == len(rows)
    assert rows == original


def test_full_selected_week_can_have_eight_delivery_dates_and_unique_long_menu_sheets():
    rows = [delivery_row(menu_id=menu, menu_name="同名超長菜單名稱" * 8, requirement_date=date(2026, 11, day),
                         meal_type_name=meal, dish_id=f"{menu}-{day}-{meal}", quantity=Decimal("1"))
            for menu in ("menu-a", "menu-b") for day in range(2, 9) for meal in ("早餐", "晚餐")]
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows, breakfast_previous_day=True)))
    assert len(workbook.sheetnames) == 16 and len(set(workbook.sheetnames)) == 16
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert "11 月 1 日（星期日）" in workbook.worksheets[0]["A3"].value
    assert "11 月 8 日（星期日）" in workbook.worksheets[-1]["A3"].value
    assert sum(sum(sheet.cell(i, 1).value is not None for i in range(6, sheet.max_row + 1)) for sheet in workbook) == len(rows)
    assert all(sheet["A6"].value == 1 for sheet in workbook)


@pytest.mark.parametrize("invalid", [None, "", "  ", 123, ["早餐"], {"name": "早餐"}])
def test_special_export_requires_reliable_meal_including_raw_snapshot_types(invalid):
    supplier_id = uuid4()
    sources = [source(uuid4(), "菜單", "2026-11-02", uuid4(), "早餐蛋餅", "1")]
    sources[0]["meal_type_name"] = invalid
    items = [item(supplier_id, sources, "1")]
    assert rows_from_snapshot_items(items, supplier_id)  # legacy ordinary export remains accepted
    with pytest.raises(SupplierWeeklyExportError, match="SUPPLIER_ORDER_MEAL_UNAVAILABLE"):
        rows_from_snapshot_items(items, supplier_id, require_meal_name=True)
    with pytest.raises(SupplierWeeklyExportError, match="SUPPLIER_ORDER_MEAL_UNAVAILABLE"):
        supplier_weekly_workbook([delivery_row(meal_type_name=invalid)], breakfast_previous_day=True)


def test_missing_historical_meal_and_non_breakfast_invalid_row_reject_whole_special_export():
    supplier_id = uuid4()
    source_row = source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜色", "1")
    del source_row["meal_type_name"]
    with pytest.raises(SupplierWeeklyExportError, match="SUPPLIER_ORDER_MEAL_UNAVAILABLE"):
        rows_from_snapshot_items([item(supplier_id, [source_row], "1")], supplier_id, require_meal_name=True)
    with pytest.raises(SupplierWeeklyExportError, match="SUPPLIER_ORDER_MEAL_UNAVAILABLE"):
        supplier_weekly_workbook([delivery_row(), delivery_row(meal_type_name=None)], breakfast_previous_day=True)
    assert mapped(SupplierWeeklyExportError("SUPPLIER_ORDER_MEAL_UNAVAILABLE")).status_code == 422


def test_special_mode_does_not_bypass_source_allocation_validation():
    supplier_id = uuid4()
    with pytest.raises(SupplierWeeklyExportError, match="SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE"):
        rows_from_snapshot_items([item(supplier_id, [source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜色")], "99")],
                                 supplier_id, require_meal_name=True)


def test_service_preserves_history_and_names_file_by_delivery_range(monkeypatch):
    from app.domains.exports.service import ExportService
    from app.domains.snapshots.service import SnapshotService

    supplier_id = uuid4()
    breakfast = source(uuid4(), "菜單", "2026-11-02", uuid4(), "菜一", "1")
    breakfast["meal_type_name"] = "早餐"
    dinner = {**breakfast, "requirement_date": "2026-11-08", "meal_type_name": "晚餐", "dish_id": str(uuid4())}
    detail = {"items": [item(supplier_id, [breakfast, dinner], "2")]}
    original = deepcopy(detail)
    monkeypatch.setattr(SnapshotService, "detail", lambda self, snapshot_id: detail)
    service = ExportService(object())  # any unexpected database query/write fails
    normal, name = service.supplier_weekly(uuid4(), supplier_id)
    assert name == "方豪_20261102-1108_叫貨單"
    explicit, explicit_name = service.supplier_weekly(uuid4(), supplier_id, False)
    assert explicit_name == name and workbook_state(explicit) == workbook_state(normal)
    shifted, name = service.supplier_weekly(uuid4(), supplier_id, True)
    assert name == "方豪_20261101-1108_叫貨單"
    assert load_workbook(BytesIO(shifted)).sheetnames == ["菜單_1101", "菜單_1108"]
    assert detail == original


def test_service_confirmed_adjustment_uses_saved_effective_lines_and_rejects_draft(monkeypatch):
    from app.domains.exports.service import ExportService
    from app.domains.order_adjustments.repository import OrderingAdjustmentRepository
    from app.domains.snapshots.service import SnapshotService

    supplier_id = uuid4()
    detail = {"ordering_adjustment_sheet_id": uuid4(), "ordering_adjustment_status": "confirmed", "items": []}
    line = SimpleNamespace(source_supplier_id=supplier_id, source_menu_id=uuid4(), menu_name_snapshot="歷史菜單",
        requirement_date=date(2026, 11, 2), meal_type_name_snapshot="員工早餐", source_dish_id=uuid4(),
        dish_name_snapshot="菜色", ingredient_name_snapshot="蔬菜", adjusted_quantity=Decimal("5.25"),
        adjusted_unit="斤", system_quantity=Decimal("4000"), system_unit="g", supplier_name_snapshot="方豪",
        meal_type_sort_order_snapshot=1, menu_dish_sort_order_snapshot=1, dish_ingredient_sort_order_snapshot=1, id=uuid4())
    original = deepcopy(vars(line))
    monkeypatch.setattr(SnapshotService, "detail", lambda self, snapshot_id: detail)
    monkeypatch.setattr(OrderingAdjustmentRepository, "lines", lambda self, sheet_id: [line])
    payload, name = ExportService(object()).supplier_weekly(uuid4(), supplier_id, True)
    sheet = load_workbook(BytesIO(payload)).active
    assert sheet.title == "歷史菜單_1101" and name == "方豪_20261101-1101_叫貨單"
    assert sheet["C6"].value == 3.15 and sheet["D6"].value == "kg"
    assert vars(line) == original
    detail["ordering_adjustment_status"] = "draft"
    with pytest.raises(SupplierWeeklyExportError, match="ADJUSTMENT_EXPORT_STATUS_INVALID"):
        ExportService(object()).supplier_weekly(uuid4(), supplier_id, True)


def test_route_boolean_default_validation_auth_and_error_without_database(monkeypatch):
    from fastapi.testclient import TestClient
    from app.api.dependencies import get_db_session
    from app.domains.auth.dependencies import get_current_user
    from app.domains.exports.service import ExportService
    from app.main import create_app

    calls = []
    def export(self, snapshot_id, supplier_id, breakfast_previous_day=False):
        calls.append(breakfast_previous_day)
        if breakfast_previous_day:
            raise SupplierWeeklyExportError("SUPPLIER_ORDER_MEAL_UNAVAILABLE")
        return b"excel-test", "叫貨單"
    monkeypatch.setattr(ExportService, "supplier_weekly", export)
    app = create_app()
    app.dependency_overrides[get_db_session] = lambda: object()
    app.dependency_overrides[get_current_user] = lambda: object()
    path = f"/api/v1/exports/requirement-snapshots/{uuid4()}/suppliers/{uuid4()}/weekly-order.xlsx"
    with TestClient(app) as client:
        assert client.get(path).content == client.get(path + "?breakfast_previous_day=false").content
        response = client.get(path + "?breakfast_previous_day=true")
        assert response.status_code == 422 and response.json()["detail"]["code"] == "SUPPLIER_ORDER_MEAL_UNAVAILABLE"
        assert client.get(path + "?breakfast_previous_day=invalid").status_code == 422
        assert calls == [False, False, True]
        del app.dependency_overrides[get_current_user]
        assert client.get(path + "?breakfast_previous_day=true").status_code == 401


def test_special_workbook_round_trip_preserves_layout_units_and_source(tmp_path):
    rows = [delivery_row(ingredient_name=f"品項-{index}", unit=unit, quantity=Decimal(quantity))
            for index, (unit, quantity) in enumerate([
                ("g", "6500"), ("斤", "5.25"), ("kg", "4.8"), ("片", "120"), ("個", "450"), ("隻", "210"),
                ("包", "10"), ("盒", "5"), ("袋", "3"), ("箱", "2"), ("罐", "3"), ("桶", "4"), ("ml", "300"), ("L", "2"),
            ])]
    original = deepcopy(rows)
    path = tmp_path / "supplier-breakfast.xlsx"
    path.write_bytes(supplier_weekly_workbook(rows, breakfast_previous_day=True))
    sheet = load_workbook(path).active
    assert sheet.title == "測試菜單_1101" and "星期日" in sheet["A3"].value
    actual = {sheet.cell(row, 2).value: (sheet.cell(row, 3).value, sheet.cell(row, 4).value) for row in range(6, 20)}
    assert actual["品項-0"] == (6.5, "kg") and actual["品項-1"] == (3.15, "kg") and actual["品項-2"] == (4.8, "kg")
    for index, row in enumerate(rows[3:], 3):
        assert actual[f"品項-{index}"] == (int(row["quantity"]), row["unit"])
    assert all(sheet.cell(row, 3).data_type == "n" for row in range(6, 20))
    assert sheet.max_row == 24 and sheet.max_column == 4
    assert sheet["A1"].alignment.horizontal == "center"
    assert all(sheet.cell(row, col).alignment.horizontal == "left" for row in range(5, 25) for col in range(1, 5))
    assert all(sheet.cell(row, col).border.bottom.style == "thin" and sheet.cell(row, col).border.bottom.color.rgb == "00000000"
               for row in range(5, 25) for col in range(1, 5))
    assert all(sheet.cell(5, col).fill.fill_type is None for col in range(1, 5))
    assert sheet.auto_filter.ref is None and not sheet.tables
    assert sheet.page_setup.paperSize == 9 and sheet.page_setup.orientation == "portrait"
    assert rows == original
