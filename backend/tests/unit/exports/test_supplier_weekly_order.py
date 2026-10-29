from collections import Counter
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import NAMESPACE_URL, uuid5, uuid4
from xml.etree.ElementTree import tostring

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.domains.exports.supplier_weekly import (
    rows_from_snapshot_items, rows_from_adjustment_lines, supplier_weekly_workbook,
)
from app.domains.requirements.calculator import calculate_requirement_rows
from app.domains.requirements.repository import RequirementRepository
from app.domains.requirements.source_order import (
    attach_source_order, saved_source_orders,
)
from app.domains.snapshots.fingerprint import normalize, content_fingerprint, result_fingerprint
from app.domains.snapshots.models import RequirementSnapshotItem
from app.domains.order_adjustments.service import OrderingAdjustmentService
from tests.unit.requirements.test_calculator import source
from tests.unit.snapshots.test_result_consistency import flow  # noqa: F401 -- shared real writer fixture


def uid(value):
    return uuid5(NAMESPACE_URL, value)


SUPPLIER = uid("order-supplier")
MENU = uid("order-menu")
COLUMNS = [{"id": str(uid(f"column-{index}")), "sort_order": index} for index in (1, 2, 3)]


def scheduled(name, *, day=date(2026, 11, 9), meal="午餐", meal_order=2,
              dish="第一菜", column=1, dish_order=1, ingredient_order=1,
              quantity="1000", unit="g", supplier=SUPPLIER):
    dish_key = f"{day}-{meal}-{dish}"
    return source(
        menu_id=MENU, menu_name="測試菜單", menu_date=day, menu_day_id=uid(f"{day}-{meal}"),
        meal_type_id=uid(meal), meal_type_name=meal, meal_type_sort_order=meal_order,
        menu_dish_id=uid(dish_key), dish_id=uid(dish), dish_name=dish,
        menu_meal_type_column_id=uid(f"column-{column}") if column is not None else None,
        menu_dish_sort_order=dish_order, _menu_columns=deepcopy(COLUMNS),
        recipe_detail_id=uid(f"{dish}-{name}"), dish_ingredient_sort_order=ingredient_order,
        ingredient_id=uid(name), ingredient_code=name, ingredient_name=name,
        recipe_quantity=Decimal(quantity), recipe_unit=unit, base_unit=unit,
        diner_count=1, loss_rate=Decimal(0), supplier_id=supplier,
        supplier_name="測試供應商", purchase_unit=unit,
    )


def example_sources():
    return [
        scheduled("早餐食材", meal="早餐", meal_order=1),
        scheduled("午餐後菜先料", dish="後菜", column=2, dish_order=1, ingredient_order=1),
        scheduled("午餐後菜後料", dish="後菜", column=2, dish_order=1, ingredient_order=2, quantity="6500"),
        scheduled("午餐前菜", dish="前菜", column=1, dish_order=9, quantity="5.25", unit="斤"),
        scheduled("晚餐食材", meal="晚餐", meal_order=3),
        scheduled("隔日早餐食材", day=date(2026, 11, 10), meal="早餐", meal_order=1),
    ]


def snapshot_items(sources):
    rows, _, _ = calculate_requirement_rows(sources)
    return [{
        "id": str(uid(row["row_key"])), "supplier_id": str(row["supplier_id"]),
        "supplier_name_snapshot": row["supplier_name"], "ingredient_name_snapshot": row["ingredient_name"],
        "requirement_unit": row["requirement_unit"], "purchase_unit_snapshot": row["suggested_purchase_unit"],
        "adjusted_quantity": row["suggested_purchase_quantity"], "source_summary": normalize(row["schedules"]),
    } for row in rows]


def exported(items, *, early=False):
    rows = rows_from_snapshot_items(items, SUPPLIER, require_meal_name=early)
    return supplier_weekly_workbook(rows, breakfast_previous_day=early)


def names(sheet):
    return [sheet.cell(row, 2).value for row in range(6, sheet.max_row + 1)
            if sheet.cell(row, 2).value is not None]


def content(book):
    return Counter((sheet.title, sheet["A3"].value, sheet["A4"].value,
                    sheet.cell(row, 2).value, sheet.cell(row, 3).value, sheet.cell(row, 4).value)
                   for sheet in book for row in range(6, sheet.max_row + 1)
                   if sheet.cell(row, 2).value is not None)


def test_repository_captures_layout_in_one_statement_not_another_read():
    calls = []
    session = SimpleNamespace(execute=lambda statement: (
        calls.append(statement) or SimpleNamespace(mappings=lambda: [scheduled("食材")])
    ))
    criteria = SimpleNamespace(menu_ids=[MENU], selected_dates=None, start_date=None)
    result = RequirementRepository(session).source_rows(criteria)
    assert len(calls) == 1
    sql = str(calls[0].compile(dialect=postgresql.dialect()))
    assert "jsonb_agg" in sql and "GROUP BY" in sql
    assert result[0]["source_order"]["visual_position"] == 1
    assert "_menu_columns" not in result[0]


def test_actual_placement_and_recipe_order_survive_source_shuffle_and_normalize():
    original = example_sources()
    original_copy = deepcopy(original)
    ordered = attach_source_order(original)
    shuffled = attach_source_order(list(reversed(original)))
    assert snapshot_items(ordered) == list(reversed(snapshot_items(shuffled)))
    assert original == original_copy
    items = snapshot_items(ordered)
    items.reverse()
    for item in items:
        item["source_summary"].reverse()
    sheet = load_workbook(BytesIO(exported(items)))["測試菜單_1109"]
    assert names(sheet) == ["早餐食材", "午餐前菜", "午餐後菜先料", "午餐後菜後料", "晚餐食材"]


def test_early_breakfast_sorts_usage_day_before_next_day_breakfast():
    items = snapshot_items(attach_source_order(example_sources()))
    book = load_workbook(BytesIO(exported(items, early=True)))
    assert names(book["測試菜單_1109"]) == ["午餐前菜", "午餐後菜先料", "午餐後菜後料", "晚餐食材", "隔日早餐食材"]
    assert names(book["測試菜單_1108"]) == ["早餐食材"]


def test_tied_meal_sort_orders_keep_each_meal_together_in_menu_api_order():
    sources = [scheduled(f"{meal}-{column}", meal=meal, meal_order=2,
                         dish=f"菜{column}", column=column, dish_order=column)
               for meal in ("餐別甲", "餐別乙") for column in (1, 2)]
    tagged = attach_source_order(list(reversed(sources)))
    expected = [row["ingredient_name"] for row in sorted(tagged, key=lambda row: (
        row["source_order"]["meal_position"], row["source_order"]["visual_position"]))]
    assert names(load_workbook(BytesIO(exported(snapshot_items(tagged)))).active) == expected


def test_fallback_empty_columns_overflow_and_supplier_filter_use_full_grid():
    sources = [scheduled("其他供應商", dish="佔位菜", column=1, dish_order=1, supplier=uid("other")),
               scheduled("固定欄", dish="固定菜", column=2, dish_order=2),
               scheduled("重複欄", dish="重複欄菜", column=1, dish_order=3),
               scheduled("未指定欄", dish="溢出菜", column=None, dish_order=4)]
    tagged = attach_source_order(sources)
    assert [row["source_order"]["visual_position"] for row in tagged] == [1, 2, 3, 4]
    sheet = load_workbook(BytesIO(exported(snapshot_items(tagged)))).active
    assert names(sheet) == ["固定欄", "重複欄", "未指定欄"]


@pytest.mark.parametrize("early", [False, True])
def test_same_actual_orders_and_full_original_format_only_sequence_changes(early):
    tagged = attach_source_order(example_sources() + [scheduled("午餐前菜", meal="晚餐", meal_order=3, quantity="2000", unit="斤")])
    items = snapshot_items(tagged)
    legacy = deepcopy(items)
    for item in legacy:
        for source_row in item["source_summary"]:
            source_row.pop("source_order", None)
    old = load_workbook(BytesIO(exported(legacy, early=early)))
    new = load_workbook(BytesIO(exported(items, early=early)))
    assert content(old) == content(new)
    assert old.sheetnames == new.sheetnames
    # Same ingredient from two dishes remains two rows, not an extra aggregate.
    assert sum(name == "午餐前菜" for sheet in new for name in names(sheet)) == 2
    for a, b in zip(old, new):
        assert list(a.values)[:5] == list(b.values)[:5]
        assert a.max_column == b.max_column == 4
        assert a.max_row == b.max_row
        assert str(a.merged_cells) == str(b.merged_cells)
        assert tostring(a.page_setup.to_tree()) == tostring(b.page_setup.to_tree())
        assert tostring(a.page_margins.to_tree()) == tostring(b.page_margins.to_tree())
        assert str(a.print_area) == str(b.print_area)
        assert a.print_title_rows == b.print_title_rows
        assert {k: v.width for k, v in a.column_dimensions.items()} == {k: v.width for k, v in b.column_dimensions.items()}
        assert not b.tables and b.auto_filter.ref is None
        for old_row, new_row in zip(a, b):
            for old_cell, new_cell in zip(old_row, new_row):
                assert old_cell._style == new_cell._style
        for row in range(6, b.max_row - 4):
            assert b.cell(row, 3).data_type == "n"
            assert all(b.cell(row, col).alignment.horizontal == "left" for col in range(1, 5))
        assert all(b.cell(row, 2).value is None for row in range(b.max_row - 4, b.max_row + 1))


@pytest.mark.parametrize("day", [date(2026, 11, 2), date(2026, 12, 1), date(2027, 1, 1)])
def test_delivery_boundaries_unchanged(day):
    items = snapshot_items(attach_source_order([scheduled("早餐", day=day, meal="早餐", meal_order=1)]))
    from datetime import timedelta
    assert load_workbook(BytesIO(exported(items, early=True))).active.title.endswith((day-timedelta(days=1)).strftime("_%m%d"))


@pytest.mark.parametrize("invalid", [None, {}, {"version": 2}, {"version": True}])
def test_legacy_or_untrusted_order_preserves_whole_sheet_legacy_behavior(invalid):
    legacy = snapshot_items(attach_source_order(example_sources()))
    for item in legacy:
        for row in item["source_summary"]:
            row.pop("source_order", None)
    changed = deepcopy(legacy)
    for item in changed:
        for row in item["source_summary"]:
            row["source_order"] = invalid
    assert list(load_workbook(BytesIO(exported(legacy))).active.values) == list(load_workbook(BytesIO(exported(changed))).active.values)


def test_mixed_sheet_does_not_partly_apply_new_sort():
    items = snapshot_items(attach_source_order(example_sources()))
    legacy = deepcopy(items)
    for item in legacy:
        for row in item["source_summary"]:
            row.pop("source_order", None)
    items[0]["source_summary"][0].pop("source_order")
    a, b = load_workbook(BytesIO(exported(items))), load_workbook(BytesIO(exported(legacy)))
    assert list(a["測試菜單_1109"].values) == list(b["測試菜單_1109"].values)


def test_frozen_sources_retain_saved_order_never_current_layout(flow):
    world = flow.world
    world.live = attach_source_order([world.row(_menu_columns=[])])
    viewed = world.calculate()
    sheet = world.capture(world.live)
    line = world.saved_lines(sheet)[0]
    item = world.items[line.snapshot_item_id]
    item.source_summary = normalize(viewed["rows"][0]["schedules"])
    history = deepcopy(item.source_summary)
    line.adjusted_quantity, line.adjusted_unit = Decimal("3.25"), "kg"
    world.live[0]["source_order"]["visual_position"] = 77
    frozen = world.calculate([sheet])
    assert frozen["rows"][0]["schedules"][0]["source_order"]["visual_position"] == 1
    assert frozen["rows"][0]["requirement_quantity"] == Decimal("3.25")
    assert item.source_summary == history
    export_rows = rows_from_adjustment_lines([line], world.supplier_id, snapshot_items=[item])
    assert export_rows[0]["source_order"]["visual_position"] == 1
    assert export_rows[0]["quantity"] == Decimal("3.25")
    item.source_summary = []
    assert "source_order" not in OrderingAdjustmentService._frozen_source_row(line, item)


def test_legacy_confirmed_source_does_not_inherit_live_order(flow):
    world = flow.world
    sheet = world.capture(world.live)
    world.live = attach_source_order([{**row, "_menu_columns": []} for row in world.live])
    assert world.live[0]["source_order"]["version"] == 1
    assert "source_order" not in world.calculate([sheet])["rows"][0]["schedules"][0]


def test_confirmed_export_service_reads_saved_order_and_effective_amount_only(flow, monkeypatch):
    from app.domains.exports.service import ExportService
    from app.domains.snapshots.service import SnapshotService
    from app.domains.order_adjustments import repository as adjustments_repository

    world = flow.world
    world.live = attach_source_order([world.row(_menu_columns=[])])
    viewed = world.calculate()
    sheet = world.capture(world.live)
    line = world.saved_lines(sheet)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal("5.25"), "斤"
    detail = {"ordering_adjustment_sheet_id": sheet.id, "ordering_adjustment_status": "confirmed",
              "items": [{"source_summary": normalize(viewed["rows"][0]["schedules"])}]}
    original = deepcopy(detail)
    monkeypatch.setattr(SnapshotService, "detail", lambda self, snapshot_id: detail)
    monkeypatch.setattr(adjustments_repository, "OrderingAdjustmentRepository", lambda session: SimpleNamespace(lines=lambda sheet_id: [line]))
    # No database or current-menu access can work through this session.
    payload, _ = ExportService(object()).supplier_weekly(sheet.baseline_snapshot_id, world.supplier_id, True)
    book = load_workbook(BytesIO(payload))
    assert book.active["C6"].value == 3.15 and book.active["D6"].value == "kg"
    assert detail == original


def test_h07_same_result_saves_explicit_order_after_normalize(flow):
    world = flow.world
    world.live = attach_source_order([world.row(_menu_columns=[])])
    viewed = world.calculate()
    flow.save(viewed)
    saved = next(x for x in flow.added if isinstance(x, RequirementSnapshotItem))
    assert saved.source_summary == normalize(viewed["rows"][0]["schedules"])
    assert saved.source_summary[0]["source_order"]["version"] == 1
    assert "commit" in flow.operations


@pytest.mark.parametrize("field", ["meal_sort_order", "meal_position", "visual_position", "dish_sort_order", "ingredient_position", "ingredient_sort_order"])
def test_h07_changed_order_rejects_before_write_without_quantity_changes(flow, field):
    world = flow.world
    world.live = attach_source_order([world.row(_menu_columns=[])])
    viewed = world.calculate()
    world.live[0]["source_order"][field] += 1
    with pytest.raises(HTTPException) as caught:
        flow.save(viewed)
    assert caught.value.status_code == 409 and caught.value.detail["code"] == "SNAPSHOT_RESULT_CHANGED"
    assert not flow.added and "commit" not in flow.operations


def test_dedup_and_h07_ignore_array_order_but_include_saved_positions(flow):
    world = flow.world
    world.live = attach_source_order([world.row(_menu_columns=[]), world.row(_menu_columns=[])])
    viewed = world.calculate()
    reordered = deepcopy(viewed)
    for row in reordered["rows"]:
        row["schedules"].reverse()
    assert content_fingerprint(viewed) == content_fingerprint(reordered)
    assert result_fingerprint(viewed) == result_fingerprint(reordered)
    legacy = deepcopy(viewed)
    for row in legacy["rows"]:
        for schedule in row["schedules"]:
            schedule.pop("source_order", None)
    old_hash = content_fingerprint(legacy)
    assert old_hash != content_fingerprint(viewed)
    history = deepcopy(legacy)
    flow.repository.by_content = lambda criteria, content: SimpleNamespace(id=uuid4()) if content == content_fingerprint(viewed) else None
    with pytest.raises(HTTPException) as caught:
        flow.save(viewed)
    assert caught.value.status_code == 409 and not flow.added
    assert legacy == history and content_fingerprint(legacy) == old_hash


def test_pre_upgrade_fingerprint_cannot_bypass_new_content_protection(flow):
    world = flow.world
    legacy_viewed = world.calculate()
    world.live = attach_source_order([{**row, "_menu_columns": []} for row in world.live])
    with pytest.raises(HTTPException) as caught:
        flow.save(legacy_viewed)
    assert caught.value.status_code == 409 and caught.value.detail["code"] == "SNAPSHOT_RESULT_CHANGED"
    assert not flow.added and "commit" not in flow.operations


def test_conflicting_saved_positions_are_not_trusted():
    items = snapshot_items(attach_source_order([scheduled("食材")]))
    other = deepcopy(items[0])
    other["source_summary"][0]["source_order"]["visual_position"] += 1
    assert saved_source_orders([items[0], other]) == {}


@pytest.mark.parametrize("sources", [None, "legacy", {}, [None, 7, "legacy"]])
def test_missing_legacy_optional_metadata_cannot_break_confirmed_lines(sources):
    assert saved_source_orders([{"source_summary": sources}, SimpleNamespace(source_summary=sources)]) == {}


def test_export_reconciliation_is_not_bypassed_by_sort_metadata():
    from app.domains.exports.exceptions import SupplierWeeklyExportError
    items = snapshot_items(attach_source_order([scheduled("食材")]))
    items[0]["adjusted_quantity"] += Decimal("1")
    with pytest.raises(SupplierWeeklyExportError) as caught:
        exported(items)
    assert caught.value.code == "SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE"
