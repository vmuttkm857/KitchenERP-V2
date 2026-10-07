from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.domains.menu_imports.exceptions import MenuImportValidationError
from app.domains.menu_imports.parser import MAX_FILE_SIZE, parse_menu_xlsx, safe_filename


def xlsx_bytes(
    rows,
    *,
    dates=None,
    headers=("餐別", "菜單欄位"),
    merges=(),
    extra_header=None,
    sheet_name="週菜單",
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    values = dates or [date(2026, 9, 28) + timedelta(days=offset) for offset in range(7)]
    sheet.append([*headers, *values])
    if extra_header is not None:
        sheet.cell(1, 10, extra_header)
    for row in rows:
        sheet.append(list(row))
    for merged_range in merges:
        sheet.merge_cells(merged_range)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_standard_dynamic_cross_month_blank_cells_and_deterministic_keys():
    payload = xlsx_bytes([
        ["護理之家", "主食", "白飯", None, "稀飯", None, None, None, None],
        ["護理之家", "素食+治飲湯", None, "紫菜湯", None, None, None, None, None],
        ["餓粥", "副菜1", "蒸蛋", None, None, None, None, None, None],
        ["餓粥", "副菜2", None, None, None, None, None, None, "青菜"],
        ["餓粥", "副菜3", None, None, None, None, None, None, None],
    ])

    first = parse_menu_xlsx(payload, "../../menu.xlsx")
    second = parse_menu_xlsx(payload, "menu.xlsx")

    assert first.start_date == date(2026, 9, 28)
    assert first.end_date == date(2026, 10, 4)
    assert first.meal_names == ("護理之家", "餓粥")
    assert first.column_count == 5
    assert [(row.meal_name, row.column_name, row.meal_sort_order, row.column_sort_order) for row in first.layout] == [
        ("護理之家", "主食", 1, 1),
        ("護理之家", "素食+治飲湯", 1, 2),
        ("餓粥", "副菜1", 2, 1),
        ("餓粥", "副菜2", 2, 2),
        ("餓粥", "副菜3", 2, 3),
    ]
    assert [line.original_import_name for line in first.lines] == ["白飯", "蒸蛋", "紫菜湯", "稀飯", "青菜"]
    assert all(line.diner_count == 1 for line in first.lines)
    assert [line.line_key for line in first.lines] == [line.line_key for line in second.lines]
    assert "餐別「餓粥」沒有任何菜色" not in first.fatal_errors


def test_merged_meal_and_column_expand_only_real_rows():
    payload = xlsx_bytes([
        ["中餐", "副菜", "菜A", None, None, None, None, None, None],
        [None, None, "菜B", None, None, None, None, None, None],
        [None, None, "菜C", None, None, None, None, None, None],
    ], merges=("A2:A4", "B2:B4"))

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert parsed.fatal_errors == ()
    assert [(line.meal_name, line.column_name) for line in parsed.lines] == [
        ("中餐", "副菜1"), ("中餐", "副菜2"), ("中餐", "副菜3"),
    ]
    assert [line.column_sort_order for line in parsed.lines] == [1, 2, 3]


def test_explicit_numbered_columns_are_preserved_without_renumbering():
    payload = xlsx_bytes([
        ["午餐", "副菜1", "菜A", None, None, None, None, None, None],
        ["午餐", "副菜2", "菜B", None, None, None, None, None, None],
        ["午餐", "副菜3", "菜C", None, None, None, None, None, None],
    ])

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert parsed.fatal_errors == ()
    assert [line.column_name for line in parsed.lines] == ["副菜1", "副菜2", "副菜3"]


def test_merged_column_expansion_collision_is_fatal():
    payload = xlsx_bytes([
        ["午餐", "副菜", "菜A", None, None, None, None, None, None],
        [None, None, "菜B", None, None, None, None, None, None],
        ["午餐", "副菜1", "菜C", None, None, None, None, None, None],
    ], merges=("A2:A3", "B2:B3"))

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert any("展開後欄位「副菜1」名稱衝突" in error for error in parsed.fatal_errors)


def test_duplicate_slot_is_fatal_and_does_not_create_a_second_line():
    payload = xlsx_bytes([
        ["午餐", "主菜", "菜A", None, None, None, None, None, None],
        ["午餐", "主菜", "菜B", None, None, None, None, None, None],
    ])

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert len(parsed.lines) == 1
    assert parsed.lines[0].original_import_name == "菜A"
    assert any("出現兩道菜" in error for error in parsed.fatal_errors)


def test_merged_dish_area_warns_and_uses_only_top_left_without_expansion():
    payload = xlsx_bytes([
        ["早餐", "主菜", "合併菜", None, None, None, None, None, None],
        ["早餐", "副菜", None, None, None, None, None, None, None],
    ], merges=("C2:C3",))

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert len(parsed.lines) == 1
    assert parsed.lines[0].original_import_name == "合併菜"
    assert any("只讀取左上角" in warning for warning in parsed.warnings)


@pytest.mark.parametrize(
    ("dates", "extra_header", "message"),
    [
        ([date(2026, 9, 1)] * 7, None, "日期不可重複"),
        ([date(2026, 9, 1) + timedelta(days=value) for value in (0, 1, 2, 4, 5, 6, 7)], None, "連續 7 天"),
        ([date(2026, 9, 1) + timedelta(days=value) for value in (0, 1, 2, 3, 4, 6, 5)], None, "遞增"),
        ([date(2026, 9, 1) + timedelta(days=value) for value in range(7)], date(2026, 9, 8), "剛好為 7 天"),
    ],
)
def test_date_validation(dates, extra_header, message):
    payload = xlsx_bytes([["早餐", "主菜", "菜A", None, None, None, None, None, None]],
                         dates=dates, extra_header=extra_header)

    with pytest.raises(MenuImportValidationError, match=message):
        parse_menu_xlsx(payload, "menu.xlsx")


def test_missing_seventh_date_is_rejected():
    payload = xlsx_bytes([["早餐", "主菜", "菜A", None, None, None, None, None]],
                         dates=[date(2026, 9, 1) + timedelta(days=value) for value in range(6)])

    with pytest.raises(MenuImportValidationError, match="七個日期欄"):
        parse_menu_xlsx(payload, "menu.xlsx")


@pytest.mark.parametrize("headers", [("餐次", "菜單欄位"), ("餐別", "欄位"), ("餐別", "忽略", "菜單欄位")])
def test_required_headers_and_adjacency(headers):
    if len(headers) == 2:
        payload = xlsx_bytes([["早餐", "主菜", "菜A", None, None, None, None, None, None]], headers=headers)
        expected = "找不到必要欄位"
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([*headers, *[date(2026, 9, 1) + timedelta(days=value) for value in range(7)]])
        output = BytesIO(); workbook.save(output); payload = output.getvalue()
        expected = "必須是相鄰欄位"
    with pytest.raises(MenuImportValidationError, match=expected):
        parse_menu_xlsx(payload, "menu.xlsx")


def test_formula_empty_structure_and_meal_without_dishes_are_reported():
    payload = xlsx_bytes([
        ["早餐", "主菜", "=1+1", None, None, None, None, None, None],
        ["午餐", "主菜", None, None, None, None, None, None, None],
        [None, "副菜", "菜A", None, None, None, None, None, None],
        ["晚餐", None, "菜B", None, None, None, None, None, None],
    ])

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert any("菜色不可使用公式" in error for error in parsed.fatal_errors)
    assert any("缺少餐別" in error for error in parsed.fatal_errors)
    assert any("缺少菜單欄位" in error for error in parsed.fatal_errors)
    assert any("餐別「早餐」沒有任何菜色" in error for error in parsed.fatal_errors)
    assert any("餐別「午餐」沒有任何菜色" in error for error in parsed.fatal_errors)


def test_newline_in_one_dish_cell_is_not_split():
    payload = xlsx_bytes([["早餐", "主菜", "菜A\n菜B", None, None, None, None, None, None]])

    parsed = parse_menu_xlsx(payload, "menu.xlsx")

    assert len(parsed.lines) == 1
    assert parsed.lines[0].original_import_name == "菜A\n菜B"
    assert parsed.lines[0].normalized_name == "菜a 菜b"


@pytest.mark.parametrize(
    ("payload", "filename", "message"),
    [
        (b"not xlsx", "menu.xls", "只允許 .xlsx"),
        (b"not xlsx", "menu.xlsx", "無法讀取 XLSX"),
        (b"", "menu.xlsx", "不可為空"),
    ],
)
def test_invalid_files(payload, filename, message):
    with pytest.raises(MenuImportValidationError, match=message):
        parse_menu_xlsx(payload, filename)


def test_filename_sanitization_handles_client_paths_cross_platform():
    assert safe_filename(r"C:\fake\path\menu.xlsx") == "menu.xlsx"
    assert safe_filename("../../menu.xlsx") == "menu.xlsx"


def test_file_and_dimension_safety_limits():
    with pytest.raises(MenuImportValidationError, match="超過 10 MB"):
        parse_menu_xlsx(b"x" * (MAX_FILE_SIZE + 1), "menu.xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["餐別", "菜單欄位", *[date(2026, 9, 1) + timedelta(days=value) for value in range(7)]])
    sheet.cell(501, 1, "超過列數")
    output = BytesIO(); workbook.save(output)
    with pytest.raises(MenuImportValidationError, match="不可超過 500 列、50 欄"):
        parse_menu_xlsx(output.getvalue(), "menu.xlsx")


def test_sheet_selection_is_explicit_when_multiple_sheets_match():
    workbook = Workbook()
    first = workbook.active
    first.title = "第一週"
    second = workbook.create_sheet("第二週")
    for sheet in (first, second):
        sheet.append(["餐別", "菜單欄位", *[date(2026, 9, 1) + timedelta(days=value) for value in range(7)]])
        sheet.append(["早餐", "主菜", "菜A", None, None, None, None, None, None])
    output = BytesIO(); workbook.save(output); payload = output.getvalue()

    with pytest.raises(MenuImportValidationError, match="請指定 sheet_name"):
        parse_menu_xlsx(payload, "menu.xlsx")
    parsed = parse_menu_xlsx(payload, "menu.xlsx", "第二週")
    assert parsed.sheet_name == "第二週"
