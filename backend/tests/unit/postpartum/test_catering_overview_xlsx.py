from copy import deepcopy
from datetime import date
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

from app.domains.postpartum.catering_overview_xlsx import render_catering_overview_xlsx


def item(index, *, preparation="要中藥", restrictions=("不牛",), note="少鹽"):
    return {
        "case_id": uuid4(), "case_number": f"C{index:03}", "name": f"個案{index}",
        "current_room": str(index), "preparation_mode": "herbal", "preparation_mode_label": preparation,
        "service_start_date": date(2026, 8, 23), "service_start_meal": "lunch", "service_start_meal_label": "午餐",
        "service_end_date": None, "service_end_meal": None, "service_end_meal_label": None,
        "restriction_groups": [
            {"id": uuid4(), "name": name, "color": "#AA0000", "is_active": True}
            for name in restrictions
        ],
        "service_note": note, "service_meals": [],
    }


def payload(items, target=date(2026, 9, 22)):
    return {"target_date": target, "weekday_label": "星期二", "total": len(items), "items": items}


def workbook_for(data, *, rich_text=False):
    return load_workbook(BytesIO(render_catering_overview_xlsx(data)), rich_text=rich_text)


def rgb(cell):
    return cell.font.color.rgb[-6:] if cell.font.color and cell.font.color.type == "rgb" else None


def test_workbook_matches_golden_structure_typography_and_print_settings():
    workbook = workbook_for(payload([
        item(501, preparation="米酒水＋麻油"),
        item(513, preparation="不中藥"),
    ]))
    sheet = workbook["1150922"]
    assert list(sheet.merged_cells.ranges)[0].coord == "A1:F1"
    assert sheet["A1"].value == "製表日期115/9/22"
    assert sheet["A1"].font.name == "標楷體" and sheet["A1"].font.sz == 16 and sheet["A1"].font.bold
    assert sheet.row_dimensions[1].height == 30
    assert [sheet.column_dimensions[column].width for column in "ABCDEF"] == [
        27.6328125, 28.36328125, 58.08984375, 29.6328125, 31.453125, 58.453125,
    ]
    assert sheet["A2"].value == "501" and sheet["D2"].value == "513"
    assert sheet["A3"].value == "個案501" and sheet["D3"].value == "個案513"
    assert sheet["B3"].value == "8/23午餐起" and sheet["E3"].value == "8/23午餐起"
    assert sheet["B2"].value == "米酒水" and sheet["E2"].value == "不中藥"
    assert sheet["A2"].font.sz == 36 and sheet["A2"].font.bold and rgb(sheet["A2"]) == "7030A0"
    assert sheet["B2"].font.sz == 28 and rgb(sheet["B2"]) == "FF0000"
    assert rgb(sheet["E2"]) == "00B0F0"
    assert sheet["A3"].font.name == "標楷體" and sheet["A3"].font.sz == 18
    assert sheet["B3"].font.name == "標楷體" and sheet["B3"].font.sz == 24
    assert sheet["C2"].font.name == "標楷體" and sheet["C2"].font.sz == 26
    assert sheet["C2"].alignment.wrap_text and sheet["C2"].alignment.vertical == "center"
    assert all(getattr(sheet["C2"].border, edge).style == "thin" for edge in ("top", "bottom", "left", "right"))
    assert sheet.row_dimensions[2].height >= 243 and sheet.row_dimensions[3].height == 35.5
    assert sheet.page_setup.orientation == "portrait" and sheet.page_setup.paperSize == 9 and sheet.page_setup.scale == 40
    assert sheet.page_margins.left == 0.3937007874015748 and sheet.page_margins.right == 0
    assert sheet.page_margins.top == 0 and sheet.page_margins.bottom == 0
    assert abs(sheet.page_margins.header - 0.11811023622047245) < 1e-12
    assert abs(sheet.page_margins.footer - 0.11811023622047245) < 1e-12
    assert sheet.sheet_view.zoomScale == 25 and sheet.max_column == 6
    assert not any(sheet.column_dimensions[column].hidden for column in "ABCDEF")


def test_pairs_cases_left_right_then_next_pair_and_keeps_odd_slot_empty():
    sheet = workbook_for(payload([item(501), item(502), item(503)]))["1150922"]
    assert [sheet[cell].value for cell in ("A2", "D2", "A4")] == ["501", "502", "503"]
    assert [sheet[cell].value for cell in ("D4", "E4", "F4", "D5", "E5", "F5")] == [None] * 6
    for cell in ("D4", "E4", "F4", "D5", "E5", "F5"):
        assert all(getattr(sheet[cell].border, edge).style == "thin" for edge in ("top", "bottom", "left", "right"))


def test_presentation_mapping_rich_text_and_source_data_are_preserved():
    items = [
        item(501, preparation="要中藥"),
        item(502, preparation="米酒水+麻油"),
        item(503, preparation="米酒水"),
        item(504, preparation="其他調理"),
    ]
    original = deepcopy(items)
    sheet = workbook_for(payload(items), rich_text=True)["1150922"]
    assert [sheet[cell].value for cell in ("B2", "E2", "B4", "E4")] == [None, "米酒水", "米酒水", "其他調理"]
    assert isinstance(sheet["C2"].value, CellRichText)
    assert str(sheet["C2"].value) == "不牛\n少鹽"
    restriction_run, note_run = sheet["C2"].value
    assert restriction_run.font.color.rgb[-6:] == "FF0000" and restriction_run.font.b is False
    assert note_run.font.color.rgb[-6:] == "1F4E79" and note_run.font.b is True
    assert items == original


def test_long_text_expands_main_row_without_changing_small_row_height():
    case = item(501, restrictions=("不牛不羊不魚不內臟不奶不花生不芝麻不芒果" * 5,), note="每餐分開包裝" * 12)
    sheet = workbook_for(payload([case]))["1150922"]
    assert sheet.row_dimensions[2].height > 321.75
    assert sheet.row_dimensions[3].height == 35.5
    assert "不牛" in sheet["C2"].value and "每餐分開包裝" in sheet["C2"].value
