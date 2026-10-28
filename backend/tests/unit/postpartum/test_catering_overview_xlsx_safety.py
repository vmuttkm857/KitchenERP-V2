"""H06: user text stays text after an actual in-memory XLSX round trip."""

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

from app.domains.exports.safety import safe_cell_text
from app.domains.postpartum.catering_overview_xlsx import render_catering_overview_xlsx
from tests.unit.postpartum.test_catering_overview_xlsx import item, payload


RISK_TEXT = ("=1+1", "=SUM(1,2)", "+1+1", "-1+1", "@SUM(1,2)")
NORMAL_TEXT = ("陳美玲", "501-A", "中文備註：少鹽、不辣", "", None)


def assert_no_formula_xml(content):
    with ZipFile(BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                sheet = ElementTree.fromstring(archive.read(name))
                assert not list(sheet.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f"))


@pytest.mark.parametrize("value", (*RISK_TEXT, *NORMAL_TEXT, r"\=1+1"))
def test_openpyxl_formula_detection_distinguishes_equals_from_other_text(value):
    workbook = Workbook()
    workbook.active["A1"] = value
    stream = BytesIO()
    workbook.save(stream)
    cell = load_workbook(BytesIO(stream.getvalue()), data_only=False).active["A1"]
    # This documents the library behavior, rather than claiming +, - or @
    # automatically become formulas in openpyxl.
    expected_formula = isinstance(value, str) and value.startswith("=")
    assert (cell.data_type == "f") is expected_formula
    assert cell.value == (None if value == "" else value)


@pytest.mark.parametrize("value", RISK_TEXT)
@pytest.mark.parametrize("field,column,row", [
    ("current_room", 1, 2),
    ("name", 1, 3),
    ("service_note", 3, 2),
    ("restriction", 3, 2),
    ("preparation_mode_label", 2, 2),
])
def test_user_text_is_guarded_in_both_halves_of_the_export(field, column, row, value):
    cases = [item(501, restrictions=(), note=""), item(502, restrictions=(), note="")]
    for case in cases:
        if field == "restriction":
            case["restriction_groups"] = [{"name": value, "is_active": True}]
        else:
            case[field] = value
    data = payload(cases)
    original = deepcopy(data)
    content = render_catering_overview_xlsx(data)
    sheet = load_workbook(BytesIO(content), data_only=False, rich_text=True).active
    for offset in (0, 3):
        cell = sheet.cell(row, column + offset)
        assert cell.data_type == "s"
        # Match the established export policy, including conservative escaping
        # of +, - and @ even though openpyxl does not itself evaluate them.
        assert cell.value == "'" + value
    assert_no_formula_xml(content)
    assert data == original


@pytest.mark.parametrize("value", RISK_TEXT)
def test_rich_text_formula_prefixes_keep_runs_colors_and_original_text(value):
    data = payload([item(501, restrictions=(value,), note=value)])
    original = deepcopy(data)
    content = render_catering_overview_xlsx(data)
    cell = load_workbook(BytesIO(content), data_only=False, rich_text=True).active["C2"]
    assert cell.data_type == "s" and isinstance(cell.value, CellRichText)
    assert str(cell.value) == value + "\n" + value
    restriction_run, note_run = cell.value
    assert isinstance(restriction_run, TextBlock) and isinstance(note_run, TextBlock)
    assert restriction_run.font.color.rgb[-6:] == "FF0000" and not restriction_run.font.b
    assert note_run.font.color.rgb[-6:] == "1F4E79" and note_run.font.b
    plain = load_workbook(BytesIO(content), data_only=False).active["C2"]
    assert plain.data_type == "s" and plain.value == value + "\n" + value
    assert_no_formula_xml(content)
    assert data == original


@pytest.mark.parametrize("value", (*RISK_TEXT, *NORMAL_TEXT, "'=1+1"))
def test_existing_safe_cell_text_contract(value):
    expected = "" if value is None else value
    if value in RISK_TEXT:
        expected = "'" + value
    assert safe_cell_text(value) == expected


@pytest.mark.parametrize("note", (None, "", "中文備註：少鹽、不辣"))
def test_normal_names_rooms_optional_notes_and_dates_are_unchanged(note):
    case = item(501, restrictions=(), note=note)
    case.update(name="陳美玲", current_room="501-A")
    data = payload([case])
    original = deepcopy(data)
    content = render_catering_overview_xlsx(data)
    sheet = load_workbook(BytesIO(content), data_only=False).active
    assert sheet["A1"].value == "製表日期115/9/22"
    assert sheet["A2"].value == "501-A" and sheet["A2"].data_type == "s"
    assert sheet["A3"].value == "陳美玲" and sheet["A3"].data_type == "s"
    assert sheet["B3"].value == "8/23午餐起" and sheet["B3"].data_type == "s"
    assert sheet["C2"].value == (note or None)
    assert data == original


def test_text_safety_does_not_change_layout_styles_or_print_settings():
    normal = [
        item(501, preparation="其他調理", restrictions=(), note="中文備註"),
        item(502, preparation="其他調理", restrictions=("不牛",), note=""),
        item(503, preparation="其他調理", restrictions=("不牛",), note="少鹽"),
    ]
    risky = deepcopy(normal)
    for case in risky:
        case.update(name="=1+1", current_room="=1+1", preparation_mode_label="=1+1")
        case["service_note"] = "=1+1" if case["service_note"] else ""
        for group in case["restriction_groups"]:
            group["name"] = "=1+1"
    before, after = [
        load_workbook(BytesIO(render_catering_overview_xlsx(payload(cases))), rich_text=True).active
        for cases in (normal, risky)
    ]
    assert before.title == after.title
    assert before["A1"].value == after["A1"].value
    assert before.max_row == after.max_row and before.max_column == after.max_column
    assert str(before.merged_cells) == str(after.merged_cells)
    assert before.print_area == after.print_area
    assert before.print_title_rows == after.print_title_rows
    assert before.print_title_cols == after.print_title_cols
    assert {key: value.width for key, value in before.column_dimensions.items()} == {
        key: value.width for key, value in after.column_dimensions.items()
    }
    assert {key: value.height for key, value in before.row_dimensions.items()} == {
        key: value.height for key, value in after.row_dimensions.items()
    }
    for attribute in ("sheet_format", "sheet_view", "page_setup", "page_margins", "print_options"):
        assert ElementTree.tostring(getattr(before, attribute).to_tree()) == ElementTree.tostring(
            getattr(after, attribute).to_tree()
        )
    for row in before:
        for cell in row:
            assert cell._style == after[cell.coordinate]._style
