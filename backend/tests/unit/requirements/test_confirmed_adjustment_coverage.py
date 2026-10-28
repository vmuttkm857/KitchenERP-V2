"""H01 regressions through the real read-model and calculator, without a database."""

from copy import deepcopy
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

import pytest

from app.domains.order_adjustments import service as adjustments
from app.domains.order_adjustments.models import OrderingAdjustmentLine
from app.domains.requirements import service as requirements
from app.domains.requirements.exceptions import RequirementAdjustmentError
from app.domains.requirements.schemas import RequirementCriteria


class SavedSources:
    def __init__(self, monkeypatch):
        self.session = object()  # Any attempted database access or write fails.
        self.menu_id, self.meal_id, self.ingredient_id, self.supplier_id = (uuid4() for _ in range(4))
        self.menu_data = {}
        self.add_menu(self.menu_id)
        self.live = []
        self.sheets, self.lines, self.snapshots, self.items = [], [], {}, {}
        repository = SimpleNamespace(
            sheets=lambda ids: [sheet for sheet in self.sheets if sheet.id in ids],
            lines_for_sheets=lambda ids: [line for line in self.lines if line.sheet_id in ids],
            snapshot_items_by_ids=lambda ids: [item for key, item in self.items.items() if key in ids],
            snapshot=self.snapshot,
        )
        monkeypatch.setattr(adjustments, "OrderingAdjustmentRepository", lambda session: repository)
        monkeypatch.setattr(adjustments, "RequirementRepository", lambda session: self)
        monkeypatch.setattr(requirements, "RequirementRepository", lambda session: self)

    def add_menu(self, menu_id):
        self.menu_data[menu_id] = {
            "menu_id": menu_id, "menu_name": "歷史菜單", "is_active": True,
            "start_date": date(2026, 11, 2), "end_date": date(2026, 11, 8),
        }

    def row(self, **changes):
        row = {
            "menu_id": self.menu_id, "menu_name": "歷史菜單", "menu_is_active": True,
            "menu_day_id": uuid4(), "menu_date": date(2026, 11, 2),
            "meal_type_id": self.meal_id, "meal_type_name": "早餐", "meal_type_sort_order": 1,
            "meal_type_is_active": True, "menu_dish_id": uuid4(), "menu_meal_type_column_id": None,
            "menu_meal_type_column_sort_order": None, "menu_dish_sort_order": 1, "diner_count": 25,
            "dish_id": uuid4(), "dish_code": "OLD", "dish_name": "炒高麗菜", "dish_is_active": True,
            "recipe_detail_id": uuid4(), "recipe_quantity": Decimal("100"), "recipe_unit": "g",
            "loss_rate": Decimal("0"), "dish_ingredient_sort_order": 1,
            "ingredient_id": self.ingredient_id, "ingredient_code": "CABBAGE", "ingredient_name": "高麗菜",
            "base_unit": "kg", "current_price": Decimal("10"), "supplier_id": self.supplier_id,
            "supplier_code": "S1", "supplier_name": "歷史供應商", "supplier_is_active": True,
            "purchase_unit": "kg", "package_size": Decimal("1"), "minimum_order_quantity": Decimal("0"),
            "ingredient_is_active": True,
        }
        return {**row, **changes}

    @staticmethod
    def selected(row, criteria):
        return row["menu_id"] in criteria.menu_ids and (
            row["menu_date"] in criteria.selected_dates if criteria.selected_dates is not None else
            criteria.start_date <= row["menu_date"] <= criteria.end_date if criteria.start_date else True
        )

    def capture(self, rows, *, status="confirmed", **dates):
        menu_ids = list(dict.fromkeys(row["menu_id"] for row in rows))
        criteria = RequirementCriteria(menu_ids=menu_ids, **dates)
        rows = [row for row in rows if self.selected(row, criteria)]
        plans = adjustments._source_plan(rows)
        snapshot = SimpleNamespace(
            id=uuid4(), snapshot_kind="ordering_adjustment",
            criteria=criteria.model_dump(mode="json", exclude_none=True),
            source_menus=[{
                **self.menu_data[menu_id], "menu_id": str(menu_id),
                "start_date": self.menu_data[menu_id]["start_date"].isoformat(),
                "end_date": self.menu_data[menu_id]["end_date"].isoformat(),
            } for menu_id in menu_ids],
        )
        sheet = SimpleNamespace(
            id=uuid4(), baseline_snapshot_id=snapshot.id, status=status,
            source_fingerprint=adjustments._source_fingerprint(criteria, plans),
        )
        self.sheets.append(sheet)
        self.snapshots[snapshot.id] = snapshot
        by_key = {}
        for plan in plans:
            item = by_key.setdefault(plan["row_key"], SimpleNamespace(id=uuid4(), supplier_code_snapshot="S1"))
            self.items[item.id] = item
            values = {key: value for key, value in plan.items() if key not in {"row_key", "_theoretical_quantity"}}
            self.lines.append(OrderingAdjustmentLine(
                id=uuid4(), sheet_id=sheet.id, snapshot_item_id=item.id,
                adjusted_quantity=None, adjusted_unit=None, **values,
            ))
        return sheet

    def saved_lines(self, sheet):
        return [line for line in self.lines if line.sheet_id == sheet.id]

    def snapshot(self, snapshot_id):
        return self.snapshots.get(snapshot_id)

    def source_rows(self, criteria):
        return [deepcopy(row) for row in self.live if self.selected(row, criteria)]

    def menus(self, menu_ids):
        return [deepcopy(self.menu_data[menu_id]) for menu_id in menu_ids]

    def calculate(self, sheets=(), **selection):
        criteria = RequirementCriteria(
            menu_ids=selection.pop("menu_ids", [self.menu_id]),
            ordering_adjustment_sheet_ids=[sheet.id for sheet in sheets] or None, **selection,
        )
        # Exercise calculate -> requirement_context -> calculate_requirement_rows.
        # Only repository I/O is replaced, not the actual source-selection logic.
        return requirements.RequirementService(self.session).calculate(criteria)


@pytest.fixture
def world(monkeypatch):
    return SavedSources(monkeypatch)


def quantities(result):
    return {row["ingredient_code"]: row["requirement_quantity"] for row in result["rows"]}


def test_confirmed_dish_replacement_uses_only_saved_source_and_manual_quantity(world):
    old = world.row()
    sheet = world.capture([old])
    line = world.saved_lines(sheet)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal("3"), "kg"
    world.live = [world.row(ingredient_id=uuid4(), ingredient_code="WATER-SPINACH", dish_name="炒空心菜")]
    assert quantities(world.calculate([sheet])) == {"CABBAGE": Decimal("3")}
    result = world.calculate([sheet])
    assert result["rows"][0]["source_count"] == 1
    assert result["rows"][0]["schedules"][0]["dish_id"] == str(old["dish_id"])


@pytest.mark.parametrize("change", ["add", "remove", "replace"])
def test_confirmed_recipe_ingredient_changes_do_not_change_frozen_demand(world, change):
    old = world.row()
    sheet = world.capture([old])
    added = {**old, "recipe_detail_id": uuid4(), "ingredient_id": uuid4(), "ingredient_code": "GARLIC"}
    world.live = [old, added] if change == "add" else [] if change == "remove" else [added]
    assert quantities(world.calculate([sheet])) == {"CABBAGE": Decimal("2.5")}


@pytest.mark.parametrize("dates", [
    {"start_date": "2026-11-02", "end_date": "2026-11-04"},
    {"selected_dates": ["2026-11-02", "2026-11-04"]},
])
def test_partial_coverage_keeps_live_gap_later_date_and_same_named_other_menu(world, dates):
    old = world.row()
    sheet = world.capture([old], **dates)
    other_menu = uuid4()
    world.add_menu(other_menu)
    world.live = [
        world.row(ingredient_code="REPLACEMENT", ingredient_id=uuid4()),
        world.row(menu_date=date(2026, 11, 5), ingredient_code="LATER", ingredient_id=uuid4()),
        world.row(menu_id=other_menu, ingredient_code="OTHER-MENU", ingredient_id=uuid4()),
    ]
    if "selected_dates" in dates:
        world.live.append(world.row(menu_date=date(2026, 11, 3), ingredient_code="GAP", ingredient_id=uuid4()))
    actual = quantities(world.calculate([sheet], menu_ids=[world.menu_id, other_menu]))
    assert actual == {code: Decimal("2.5") for code in ["CABBAGE", "LATER", "OTHER-MENU"] + (["GAP"] if "selected_dates" in dates else [])}


def test_all_dates_uses_saved_menu_bounds_and_covers_new_meals_and_empty_original_slots(world):
    sheet = world.capture([world.row()])  # No explicit dates: saved menu range, all meals.
    world.menu_data[world.menu_id]["end_date"] = date(2026, 11, 10)
    world.live = [
        world.row(meal_type_id=uuid4(), meal_type_name="新增餐別", ingredient_code="NEW-MEAL"),
        world.row(menu_date=date(2026, 11, 4), ingredient_code="PREVIOUSLY-EMPTY"),
        world.row(menu_date=date(2026, 11, 9), ingredient_code="EXTENDED", ingredient_id=uuid4()),
    ]
    assert quantities(world.calculate([sheet])) == {"CABBAGE": Decimal("2.5"), "EXTENDED": Decimal("2.5")}


def test_request_dates_intersect_coverage_without_readding_other_frozen_dates(world):
    rows = [world.row(menu_date=date(2026, 11, day)) for day in (2, 3)]
    sheet = world.capture(rows)
    world.live = [world.row(menu_date=date(2026, 11, 3), ingredient_code="NEW")]
    result = world.calculate([sheet], selected_dates=["2026-11-03"])
    assert quantities(result) == {"CABBAGE": Decimal("2.5")}
    assert {row["requirement_date"] for row in result["daily_rows"]} == {date(2026, 11, 3)}


def test_requested_previously_empty_covered_date_does_not_accept_new_live_sources(world):
    sheet = world.capture([world.row()])
    world.live = [world.row(menu_date=date(2026, 11, 4), ingredient_code="NEW")]
    result = world.calculate([sheet], selected_dates=["2026-11-04"])
    assert result["rows"] == [] and result["daily_rows"] == []


def test_moving_same_source_id_outside_confirmed_dates_keeps_theoretical_live_quantity(world):
    old = world.row()
    sheet = world.capture([old], selected_dates=["2026-11-02"])
    line = world.saved_lines(sheet)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal("9"), "kg"
    world.live = [{**old, "menu_date": date(2026, 11, 5), "menu_day_id": uuid4()}]
    result = world.calculate([sheet])
    assert {row["requirement_date"]: row["quantity"] for row in result["daily_rows"]} == {
        date(2026, 11, 2): Decimal("9"), date(2026, 11, 5): Decimal("2.5"),
    }


def test_nonoverlapping_confirmed_and_draft_sheets_preserve_both_saved_quantities(world):
    old = world.row()
    draft_row = world.row(menu_date=date(2026, 11, 5))
    confirmed = world.capture([old], selected_dates=["2026-11-02"])
    draft = world.capture([draft_row], status="draft", selected_dates=["2026-11-05"])
    line = world.saved_lines(draft)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal("4"), "kg"
    world.live = [world.row(ingredient_code="NEW"), draft_row]
    assert quantities(world.calculate([confirmed, draft])) == {"CABBAGE": Decimal("6.5")}


@pytest.mark.parametrize("status", ["confirmed", "draft"])
def test_overlapping_coverage_with_different_source_keys_fails_closed(world, status):
    confirmed = world.capture([world.row()])
    current = world.row(ingredient_id=uuid4(), ingredient_code="NEW")
    other = world.capture([current], status=status)
    world.live = [current]
    with pytest.raises(RequirementAdjustmentError) as caught:
        world.calculate([confirmed, other])
    assert caught.value.code == "ADJUSTMENT_COVERAGE_CONFLICT"


def test_duplicate_override_error_is_preserved(world):
    old = world.row()
    first, second = world.capture([old]), world.capture([old])
    world.live = [old]
    with pytest.raises(RequirementAdjustmentError) as caught:
        world.calculate([first, second])
    assert caught.value.code == "ADJUSTMENT_DUPLICATE_OVERRIDE"


def test_two_confirmed_sheets_only_conflict_inside_requested_dates(world):
    first = world.capture([world.row()], start_date="2026-11-02", end_date="2026-11-04")
    second = world.capture([world.row(menu_date=date(2026, 11, 5))], start_date="2026-11-04", end_date="2026-11-06")
    world.live = [world.row(ingredient_code="NEW"), world.row(menu_date=date(2026, 11, 5), ingredient_code="NEW")]
    assert quantities(world.calculate([first, second], selected_dates=["2026-11-02", "2026-11-05"])) == {"CABBAGE": Decimal("5")}


@pytest.mark.parametrize("quantity,unit,expected", [("5250", "g", "5.25"), ("5.25", "斤", "3.15"), ("0", "kg", "0"), ("2.475", "kg", "2.475")])
def test_effective_manual_quantity_is_not_multiplied_or_converted_twice(world, quantity, unit, expected):
    sheet = world.capture([world.row(diner_count=400)])
    line = world.saved_lines(sheet)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal(quantity), unit
    world.live = [world.row(ingredient_code="NEW")]
    assert quantities(world.calculate([sheet])) == {"CABBAGE": Decimal(expected)}
    assert (line.adjusted_quantity, line.adjusted_unit) == (Decimal(quantity), unit)


@pytest.mark.parametrize("corruption", ["missing", "kind", "dates", "menus", "unknown_filter", "criteria", "line_scope", "item"])
def test_unknown_or_damaged_saved_coverage_fails_closed(world, corruption):
    sheet = world.capture([world.row()])
    snapshot = world.snapshots[sheet.baseline_snapshot_id]
    if corruption == "missing": del world.snapshots[snapshot.id]
    elif corruption == "kind": snapshot.snapshot_kind = "standard"
    elif corruption == "dates": del snapshot.source_menus[0]["end_date"]
    elif corruption == "menus": snapshot.source_menus[0]["menu_id"] = str(uuid4())
    elif corruption == "unknown_filter": snapshot.criteria["meal_type_ids"] = [str(world.meal_id)]
    elif corruption == "criteria": snapshot.criteria = {}
    elif corruption == "line_scope": world.saved_lines(sheet)[0].requirement_date = date(2026, 12, 1)
    elif corruption == "item": world.items.clear()
    world.live = []
    with pytest.raises(RequirementAdjustmentError) as caught:
        world.calculate([sheet])
    assert caught.value.code == "ADJUSTMENT_COVERAGE_UNAVAILABLE"


def test_draft_stale_validation_and_theoretical_path_stay_unchanged(world):
    old = world.row()
    sheet = world.capture([old], status="draft")
    world.live = [world.row(ingredient_code="NEW")]
    assert quantities(world.calculate()) == {"NEW": Decimal("2.5")}
    with pytest.raises(RequirementAdjustmentError) as caught:
        world.calculate([sheet])
    assert caught.value.code == "ADJUSTMENT_STALE"


def test_same_ingredient_different_dishes_keeps_frozen_source_rows_and_history_unchanged(world):
    rows = [world.row(), world.row(dish_name="高麗菜湯", diner_count=10)]
    sheet = world.capture(rows)
    before = deepcopy(world.snapshots[sheet.baseline_snapshot_id].__dict__)
    saved = [(line.id, line.source_line_key, line.adjusted_quantity, line.adjusted_unit) for line in world.lines]
    world.live = [world.row(ingredient_id=uuid4(), ingredient_code="NEW")]
    result = world.calculate([sheet])
    assert quantities(result) == {"CABBAGE": Decimal("3.5")}
    assert result["rows"][0]["source_count"] == 2
    assert len(result["rows"][0]["schedules"]) == 2
    assert before == world.snapshots[sheet.baseline_snapshot_id].__dict__
    assert saved == [(line.id, line.source_line_key, line.adjusted_quantity, line.adjusted_unit) for line in world.lines]


def test_new_missing_recipe_inside_confirmed_coverage_is_not_a_live_source(world):
    sheet = world.capture([world.row()])
    world.live = [world.row(recipe_detail_id=None, ingredient_id=None)]
    result = world.calculate([sheet])
    assert quantities(result) == {"CABBAGE": Decimal("2.5")}
    assert result["anomalies"] == []


def test_same_dish_in_different_meals_keeps_independent_manual_quantities(world):
    breakfast = world.row()
    lunch = {**breakfast, "menu_day_id": uuid4(), "menu_dish_id": uuid4(),
             "meal_type_id": uuid4(), "meal_type_name": "午餐", "meal_type_sort_order": 2}
    sheet = world.capture([breakfast, lunch])
    for line, quantity in zip(world.saved_lines(sheet), ("3", "5")):
        line.adjusted_quantity, line.adjusted_unit = Decimal(quantity), "kg"
    world.live = [world.row(ingredient_code="NEW")]
    result = world.calculate([sheet])
    assert quantities(result) == {"CABBAGE": Decimal("8")}
    assert len(result["rows"][0]["schedules"]) == 2
    assert {source["meal_type_name"] for source in result["rows"][0]["schedules"]} == {"早餐", "午餐"}


def test_corrected_result_flows_through_snapshot_build_supplier_excel_and_purchase_values(world):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.domains.exports.exceptions import SupplierWeeklyExportError
    from app.domains.exports.supplier_weekly import rows_from_snapshot_items, supplier_weekly_workbook
    from app.domains.purchases.calculator import purchase_values
    from app.domains.snapshots.service import SnapshotService

    sheet = world.capture([world.row(), world.row(dish_name="高麗菜湯")])
    line = world.saved_lines(sheet)[0]
    line.adjusted_quantity, line.adjusted_unit = Decimal("5.25"), "斤"
    history_before = deepcopy(world.snapshots[sheet.baseline_snapshot_id].__dict__)
    world.live = [world.row(ingredient_id=uuid4(), ingredient_code="NEW")]
    result = world.calculate([sheet])
    added, audits = [], []
    service = SnapshotService(SimpleNamespace(flush=lambda: None))
    service.repository = SimpleNamespace(
        lock_creation_key=lambda value: None, by_content=lambda *args: None,
        next_revision=lambda value: 2, add=added.append,
    )
    service.audit = SimpleNamespace(record=lambda **kwargs: audits.append(kwargs))
    header, items, _ = service.build(result["criteria"], uuid4(), result=result)
    assert header.snapshot_kind == "standard" and header.revision == 2
    assert len(added) == 2 and len(audits) == 1
    assert len(items) == 1 and items[0].requirement_quantity == Decimal("5.65")
    assert len(items[0].source_summary) == 2
    assert purchase_values(items[0])[:2] == (Decimal("5.65"), Decimal("56.5"))
    item = {column.name: getattr(items[0], column.name) for column in items[0].__table__.columns}
    rows = rows_from_snapshot_items([item], world.supplier_id, require_meal_name=True)
    workbook = load_workbook(BytesIO(supplier_weekly_workbook(rows, breakfast_previous_day=True)))
    assert len(workbook.sheetnames) == 1 and workbook.sheetnames[0].endswith("_1101")
    page = workbook.active
    assert [page.cell(5, column).value for column in range(1, 5)] == ["項次", "訂購品名", "數量", "單位"]
    assert sorted(page.cell(row, 3).value for row in (6, 7)) == [2.5, 3.15]
    assert all(page.cell(row, 3).data_type == "n" and page.cell(row, 4).value == "kg" for row in (6, 7))
    assert page.max_row == 12 and page.page_setup.orientation == "portrait" and page.page_setup.paperSize == 9
    assert all(page.cell(row, column).alignment.horizontal == "left" for row in range(5, 13) for column in range(1, 5))
    assert page.auto_filter.ref is None and not page.tables
    assert history_before == world.snapshots[sheet.baseline_snapshot_id].__dict__
    item["adjusted_quantity"] += Decimal("1")
    with pytest.raises(SupplierWeeklyExportError, match="SNAPSHOT_SOURCE_ALLOCATION_UNAVAILABLE"):
        rows_from_snapshot_items([item], world.supplier_id, require_meal_name=True)
