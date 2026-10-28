"""Real PostgreSQL H01 workflow tests; guarded by the existing database opt-in."""

from copy import deepcopy
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.api.order_adjustments.test_order_adjustment_sync_preview_api import _editor_payload, _new_dish
from tests.api.order_adjustments.test_order_adjustments_api import create_sheet
from tests.api.requirements.test_requirements_api import auth, fixture
from tests.menu_editor import save_editor
from tests.snapshot_comparison import assert_snapshot_unchanged


def _workbook_values(payload):
    return [(sheet.title, list(sheet.values)) for sheet in load_workbook(BytesIO(payload))]


@pytest.mark.parametrize("change", ["dish_replacement", "recipe_add", "recipe_remove"])
def test_confirmed_workflow_keeps_frozen_requirements_snapshot_and_supplier_excel(client, db_session, change):
    headers = auth(client, db_session)
    menu, _, dishes, ingredients, suppliers = fixture(client, headers)
    sheet = create_sheet(client, headers, menu).json()
    line = next(line for line in sheet["lines"] if line["ingredient_code_snapshot"] == "REQ-I1")
    saved = client.patch(f"/api/v1/order-adjustments/{sheet['id']}/lines", headers=headers, json={
        "lock_version": sheet["lock_version"],
        "lines": [{"id": line["id"], "adjusted_quantity": "5.25", "adjusted_unit": "斤"}],
    })
    assert saved.status_code == 200, saved.text
    confirmed = client.post(f"/api/v1/order-adjustments/{sheet['id']}/confirm", headers=headers,
                            json={"lock_version": saved.json()["lock_version"]})
    assert confirmed.status_code == 200, confirmed.text
    criteria = {"menu_ids": [menu["id"]], "selected_dates": ["2026-09-01"],
                "ordering_adjustment_sheet_ids": [sheet["id"]]}
    original = client.post("/api/v1/requirements/calculate", headers=headers, json=criteria)
    assert original.status_code == 200, original.text
    snapshot = client.post("/api/v1/requirement-snapshots", headers=headers, json={"criteria": criteria, "expected_result_fingerprint": original.json()["result_fingerprint"]})
    assert snapshot.status_code == 201, snapshot.text
    snapshot_a = deepcopy(snapshot.json())
    export_url = f"/api/v1/exports/requirement-snapshots/{snapshot_a['id']}/suppliers/{suppliers[0]['id']}/weekly-order.xlsx?breakfast_previous_day=true"
    original_excel = client.get(export_url, headers=headers)
    assert original_excel.status_code == 200, original_excel.text

    if change == "dish_replacement":
        replacement = _new_dish(client, headers, code="H01-NEW", name="新菜色", ingredient_id=ingredients[2]["id"], unit="L")
        aggregate = client.get(f"/api/v1/menus/{menu['id']}/editor", headers=headers).json()
        changed = save_editor(client, f"/api/v1/menus/{menu['id']}/editor", headers=headers,
                             json=_editor_payload(aggregate, {line["source_menu_dish_id"]: replacement["id"]}))
    else:
        recipe = client.get(f"/api/v1/dishes/{dishes[0]['id']}/recipe", headers=headers).json()
        items = [{key: item.get(key) for key in ("id", "ingredient_id", "quantity", "unit", "loss_rate", "sort_order", "notes")}
                 for item in recipe["items"]]
        if change == "recipe_add":
            items.append({"ingredient_id": ingredients[2]["id"], "quantity": "1", "unit": "L", "sort_order": 3})
        else:
            items = [item for item in items if item["ingredient_id"] != ingredients[0]["id"]]
        changed = client.put(f"/api/v1/dishes/{dishes[0]['id']}/recipe", headers=headers, json={"items": items})
    assert changed.status_code == 200, changed.text

    current = client.post("/api/v1/requirements/calculate", headers=headers, json=criteria)
    assert current.status_code == 200, current.text
    assert current.json() == original.json()
    historical = client.get(f"/api/v1/requirement-snapshots/{snapshot_a['id']}", headers=headers)
    assert historical.status_code == 200, historical.text
    assert_snapshot_unchanged(historical.json(), snapshot_a)
    after_excel = client.get(export_url, headers=headers)
    assert after_excel.status_code == 200, after_excel.text
    assert _workbook_values(after_excel.content) == _workbook_values(original_excel.content)
    current_sheet = client.get(f"/api/v1/order-adjustments/{sheet['id']}", headers=headers).json()
    assert current_sheet == confirmed.json()


def test_confirmed_partial_dates_keep_new_live_dish_outside_coverage(client, db_session):
    headers = auth(client, db_session)
    menu, _, _, ingredients, _ = fixture(client, headers)
    sheet = create_sheet(client, headers, menu).json()  # Only September 1 is covered.
    confirmed = client.post(f"/api/v1/order-adjustments/{sheet['id']}/confirm", headers=headers,
                            json={"lock_version": sheet["lock_version"]})
    assert confirmed.status_code == 200, confirmed.text
    replacement = _new_dish(client, headers, code="H01-LATER", name="未覆蓋日期新菜", ingredient_id=ingredients[0]["id"], quantity="300")
    aggregate = client.get(f"/api/v1/menus/{menu['id']}/editor", headers=headers).json()
    replacements = {dish["id"]: replacement["id"] for slot in aggregate["slots"] if slot["menu_date"] == "2026-09-03"
                    for dish in slot["dishes"] if dish["dish_id"] == sheet["lines"][0]["source_dish_id"]}
    changed = save_editor(client, f"/api/v1/menus/{menu['id']}/editor", headers=headers, json=_editor_payload(aggregate, replacements))
    assert changed.status_code == 200, changed.text
    result = client.post("/api/v1/requirements/calculate", headers=headers,
                         json={"menu_ids": [menu["id"]], "ordering_adjustment_sheet_ids": [sheet["id"]]})
    assert result.status_code == 200, result.text
    rows = [row for row in result.json()["daily_rows"] if row["ingredient_code"] == "REQ-I1"]
    assert Decimal(next(row["quantity"] for row in rows if row["requirement_date"] == "2026-09-01")) == Decimal("6.3")
    assert Decimal(next(row["quantity"] for row in rows if row["requirement_date"] == "2026-09-03")) == Decimal("12")
