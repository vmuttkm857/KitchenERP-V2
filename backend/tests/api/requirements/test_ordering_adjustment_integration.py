import uuid
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.domains.menus.models import MenuDish
from app.domains.recipes.models import DishIngredient
from app.domains.order_adjustments.models import OrderingAdjustmentLine
from app.domains.snapshots.models import RequirementSnapshotItem
from tests.api.order_adjustments.test_order_adjustments_api import create_sheet
from tests.api.requirements.test_requirements_api import auth, fixture


def _criteria(menu, sheet_id=None, **dates):
    value={"menu_ids":[menu["id"]],**dates}
    if sheet_id:value["ordering_adjustment_sheet_ids"]=[sheet_id]
    return value


def _row(body, code):
    return next(row for row in body["rows"] if row["ingredient_code"]==code)


def test_explicit_draft_applies_source_lines_and_zero_while_default_stays_theoretical(client,db_session):
    headers=auth(client,db_session);menu,*_=fixture(client,headers);sheet=create_sheet(client,headers,menu).json()
    chicken=[line for line in sheet["lines"] if line["ingredient_code_snapshot"]=="REQ-I1"]
    updated=client.patch(f"/api/v1/order-adjustments/{sheet['id']}/lines",headers=headers,json={"lock_version":1,"lines":[{"id":chicken[0]["id"],"adjusted_quantity":"0"},{"id":chicken[1]["id"],"adjusted_quantity":"2.123456"}]})
    assert updated.status_code==200,updated.text
    theoretical=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,selected_dates=["2026-09-01"])).json()
    adjusted=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,sheet["id"],selected_dates=["2026-09-01"]))
    assert adjusted.status_code==200,adjusted.text
    expected=sum((Decimal(line["adjusted_quantity"] if line["adjusted_quantity"] is not None else line["system_quantity"]) for line in updated.json()["lines"] if line["ingredient_code_snapshot"]=="REQ-I1"),Decimal("0"))
    assert Decimal(_row(theoretical,"REQ-I1")["requirement_quantity"])==Decimal("6.3")
    assert Decimal(_row(adjusted.json(),"REQ-I1")["requirement_quantity"])==expected
    assert Decimal(_row(adjusted.json(),"REQ-I1")["requirement_quantity"])!=Decimal(_row(theoretical,"REQ-I1")["requirement_quantity"])


def test_partial_date_and_multi_menu_daily_identity_are_preserved(client,db_session):
    headers=auth(client,db_session);menu_a,_,dishes,*_=fixture(client,headers)
    sheet_a=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu_a["id"]],"start_date":"2026-09-01","end_date":"2026-09-03"}}).json()
    target=next(line for line in sheet_a["lines"] if line["ingredient_code_snapshot"]=="REQ-I1" and line["requirement_date"]=="2026-09-02")
    client.patch(f"/api/v1/order-adjustments/{sheet_a['id']}/lines",headers=headers,json={"lock_version":1,"lines":[{"id":target["id"],"adjusted_quantity":"1.234567"}]})
    menu_b=client.post("/api/v1/menus",headers=headers,json={"name":"第二份菜單","start_date":"2026-09-02","end_date":"2026-09-02"}).json()
    meal=client.post(f"/api/v1/menus/{menu_b['id']}/meal-types",headers=headers,json={"name":"午餐","sort_order":1}).json()
    client.put(f"/api/v1/menus/{menu_b['id']}/editor",headers=headers,json={"slots":[{"menu_date":"2026-09-02","menu_meal_type_id":meal["id"],"dishes":[{"dish_id":dishes[0]["id"],"diner_count":10,"sort_order":1}]}]})
    response=client.post("/api/v1/requirements/calculate",headers=headers,json={"menu_ids":[menu_a["id"],menu_b["id"]],"start_date":"2026-09-02","end_date":"2026-09-02","ordering_adjustment_sheet_ids":[sheet_a["id"]]})
    assert response.status_code==200,response.text
    chicken=[row for row in response.json()["daily_rows"] if row["ingredient_code"]=="REQ-I1"]
    assert {row["menu_id"] for row in chicken}=={menu_a["id"],menu_b["id"]}
    assert Decimal(next(row["quantity"] for row in chicken if row["menu_id"]==menu_b["id"]))==Decimal("1.1")


def test_confirmed_is_frozen_while_cancelled_stale_unit_and_duplicate_overrides_are_structured(client,db_session):
    headers=auth(client,db_session);menu,*_=fixture(client,headers)
    confirmed=create_sheet(client,headers,menu).json()
    target=next(line for line in confirmed["lines"] if line["ingredient_code_snapshot"]=="REQ-I1")
    confirmed=client.patch(f"/api/v1/order-adjustments/{confirmed['id']}/lines",headers=headers,json={"lock_version":1,"lines":[{"id":target["id"],"adjusted_quantity":"0"}]}).json()
    confirmed=client.post(f"/api/v1/order-adjustments/{confirmed['id']}/confirm",headers=headers,json={"lock_version":confirmed["lock_version"]}).json()
    expected=sum((Decimal(line["adjusted_quantity"] if line["adjusted_quantity"] is not None else line["system_quantity"]) for line in confirmed["lines"] if line["ingredient_code_snapshot"]=="REQ-I1"),Decimal("0"))
    removed_recipe=db_session.get(DishIngredient,target["source_dish_ingredient_id"]);db_session.delete(removed_recipe);db_session.commit()
    response=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,confirmed["id"],selected_dates=["2026-09-01"]))
    assert response.status_code==200,response.text
    assert Decimal(_row(response.json(),"REQ-I1")["requirement_quantity"])==expected

    source=db_session.get(MenuDish,confirmed["lines"][0]["source_menu_dish_id"]);source.diner_count+=1;db_session.commit()
    stale=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu["id"]],"selected_dates":["2026-09-01"]}}).json()
    source.diner_count+=1;db_session.commit()
    response=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,stale["id"],selected_dates=["2026-09-01"]))
    assert response.status_code==409 and response.json()["detail"]["code"]=="ADJUSTMENT_STALE"

    source.diner_count-=1;db_session.commit()
    unit_sheet=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu["id"]],"selected_dates":["2026-09-02"]}}).json()
    line=db_session.get(OrderingAdjustmentLine,unit_sheet["lines"][0]["id"]);line.system_unit="斤";db_session.commit()
    response=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,unit_sheet["id"],selected_dates=["2026-09-02"]))
    assert response.status_code==409 and response.json()["detail"]["code"]=="ADJUSTMENT_STALE"

    first=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu["id"]],"selected_dates":["2026-09-03"]}}).json()
    second=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu["id"]],"start_date":"2026-09-02","end_date":"2026-09-03"}}).json()
    response=client.post("/api/v1/requirements/calculate",headers=headers,json={"menu_ids":[menu["id"]],"selected_dates":["2026-09-03"],"ordering_adjustment_sheet_ids":[first["id"],second["id"]]})
    assert response.status_code==409 and response.json()["detail"]["code"]=="ADJUSTMENT_DUPLICATE_OVERRIDE"

    cancelled=client.post(f"/api/v1/order-adjustments/{unit_sheet['id']}/cancel",headers=headers,json={"lock_version":unit_sheet["lock_version"]}).json()
    response=client.post("/api/v1/requirements/calculate",headers=headers,json=_criteria(menu,cancelled["id"],selected_dates=["2026-09-02"]))
    assert response.status_code==409 and response.json()["detail"]["code"]=="ADJUSTMENT_STATUS_INVALID"


def test_adjusted_requirements_flow_into_excel_and_standard_snapshot(client,db_session):
    headers=auth(client,db_session);menu,*_=fixture(client,headers);sheet=create_sheet(client,headers,menu).json()
    chicken=next(line for line in sheet["lines"] if line["ingredient_code_snapshot"]=="REQ-I1")
    updated=client.patch(f"/api/v1/order-adjustments/{sheet['id']}/lines",headers=headers,json={"lock_version":1,"lines":[{"id":chicken["id"],"adjusted_quantity":"0"}]})
    assert updated.status_code==200
    criteria=_criteria(menu,sheet["id"],selected_dates=["2026-09-01"])
    result=client.post("/api/v1/requirements/calculate",headers=headers,json=criteria).json();expected=Decimal(_row(result,"REQ-I1")["requirement_quantity"])
    exported=client.post("/api/v1/exports/requirements/xlsx",headers=headers,json=criteria)
    assert exported.status_code==200,exported.text
    sheet_values=load_workbook(BytesIO(exported.content))["需求彙總"]
    excel_quantity=next(Decimal(str(row[3].value)) for row in sheet_values.iter_rows(min_row=2) if row[0].value=="REQ-I1")
    assert excel_quantity==expected
    snapshot=client.post("/api/v1/requirement-snapshots",headers=headers,json={"criteria":criteria})
    assert snapshot.status_code==201,snapshot.text
    item=next(item for item in snapshot.json()["items"] if item["ingredient_code_snapshot"]=="REQ-I1")
    assert Decimal(item["requirement_quantity"])==expected


def test_adjustment_creation_ignores_nested_adjustment_context_for_theoretical_baseline(client,db_session):
    headers=auth(client,db_session);menu,*_=fixture(client,headers);first=create_sheet(client,headers,menu).json()
    client.post(f"/api/v1/order-adjustments/{first['id']}/cancel",headers=headers,json={"lock_version":1})
    response=client.post("/api/v1/order-adjustments",headers=headers,json={"criteria":{"menu_ids":[menu["id"]],"selected_dates":["2026-09-01"],"ordering_adjustment_sheet_ids":[str(uuid.uuid4())]}})
    assert response.status_code==201,response.text
    assert "ordering_adjustment_sheet_ids" not in response.json()["criteria"]
