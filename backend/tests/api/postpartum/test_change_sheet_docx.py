from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from sqlalchemy import delete

from app.domains.menus.models import MenuDish
from app.domains.postpartum.change_sheet_docx import render_daily_change_sheet_docx
from app.domains.postpartum.meals import MEAL_VALUES
from tests.api.postpartum.test_change_sheet_api import _configured_handlings
from tests.api.postpartum.test_menu_conflicts_api import TARGET_DATE, auth


MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _path(target_date=TARGET_DATE):
    return f"/api/v1/postpartum/change-sheet/daily.docx?target_date={target_date}"


def _document(response):
    assert response.content.startswith(b"PK")
    with ZipFile(BytesIO(response.content)) as archive:
        assert archive.testzip() is None
        assert "[Content_Types].xml" in archive.namelist()
        assert "word/document.xml" in archive.namelist()
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(name))
    return Document(BytesIO(response.content))


def _all_text(document):
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


def _minimal_daily(meal):
    empty = {
        "summary": {
            "replacement_group_count": 0, "replacement_item_count": 0,
            "manual_acknowledgement_count": 0, "requires_reconfirmation_count": 0,
        },
        "replacement_groups": [], "manual_acknowledgements": [],
        "requires_reconfirmation": [], "warnings": [], "has_changes": False,
    }
    labels = ("早餐", "早點", "午餐", "午點", "晚餐", "晚點")
    meals = [{**empty, "meal_label": label} for label in labels]
    meals[2] = meal
    return {
        "target_date": TARGET_DATE,
        "summary": meal["summary"],
        "warnings": meal.get("warnings", []),
        "meals": meals,
    }


def test_empty_daily_docx_is_authenticated_valid_and_keeps_canonical_meal_order(client, db_session):
    headers = auth(client, db_session)
    response = client.get(_path(), headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == MIME
    assert "attachment" in response.headers["content-disposition"]
    assert TARGET_DATE in response.headers["content-disposition"]
    document = _document(response)
    text = _all_text(document)
    assert "月子餐每日異動單" in text
    assert "日期：2026/09/10（星期四）" in text
    positions = [text.index(label) for label in ("早餐", "早點", "午餐", "午點", "晚餐", "晚點")]
    assert positions == sorted(positions)
    assert text.count("無異動") == len(MEAL_VALUES)
    assert client.get(_path("not-a-date"), headers=headers).status_code == 422
    assert client.get(_path()).status_code == 401


def test_daily_docx_uses_daily_read_model_quantity_items_rooms_restrictions_and_acknowledgement(client, db_session):
    headers, first_case, second_case, _, _, group, _ = _configured_handlings(client, db_session)
    response = client.get(_path(), headers=headers)
    assert response.status_code == 200, response.text
    document = _document(response)
    text = _all_text(document)
    assert group["replacement_dish"]["name"] in text
    assert group["replacement_dish"]["code"] not in text
    assert first_case["current_room"] in text and second_case["current_room"] in text
    assert "不牛" in text or "直接菜色" in text
    assert "人工確認，不需替代" in text
    assert "人工確認備註" in text
    assert "廚房製作" not in text
    assert not any(table.rows[0].cells[0].text == "替代菜" for table in document.tables)
    details = next(table for table in document.tables if table.rows[0].cells[0].text == "床號" and table.rows[0].cells[3].text == "替代菜")
    assert all("\n" not in row.cells[1].text for row in details.rows[1:])
    assert len(details.rows) == 4  # header + all three handling items; no dedupe
    assert sum(1 for row in details.rows[1:] if group["replacement_dish"]["name"] in row.cells[3].text) == 3
    raw_ids = [group["id"], first_case["id"], second_case["id"]]
    assert all(value not in text for value in raw_ids)


def test_normal_docx_omits_coverage_diagnostics_but_keeps_operational_content():
    diagnostic = "禁忌群組『不內臟(含腰子)』無法由菜名或食材關聯自動判定"
    group = {
        "replacement_dish": {"code": "09", "name": "咕咾肉(玖)"},
        "quantity": 2, "case_rooms": ["505", "508"], "note": "少油",
        "status": "replaced", "warnings": [{"code": "COVERAGE", "message": diagnostic}],
        "items": [{
            "current_room": "505", "case_name": "王小姐", "case_number": "P-001",
            "original_dish": {"code": "06", "name": "無骨雞排"},
            "replacement_dish": {"code": "09", "name": "咕咾肉(玖)"},
            "restriction_groups": [{"name": "不內臟"}], "status": "replaced",
            "warnings": [{"code": "COVERAGE", "message": diagnostic}],
        }],
    }
    meal = {
        "meal_label": "午餐", "has_changes": True,
        "summary": {
            "replacement_group_count": 1, "replacement_item_count": 1,
            "manual_acknowledgement_count": 0, "requires_reconfirmation_count": 0,
        },
        "replacement_groups": [group], "manual_acknowledgements": [],
        "requires_reconfirmation": [], "warnings": [{"code": "COVERAGE", "message": diagnostic}],
    }
    document = Document(BytesIO(render_daily_change_sheet_docx(_minimal_daily(meal))))
    text = _all_text(document)
    assert diagnostic not in text
    for expected in ("咕咾肉(玖)", "505", "王小姐", "無骨雞排", "不內臟"):
        assert expected in text
    for removed in (
        "廚房製作", "替代製作：", "人工確認：", "待重新確認：",
        "P-001", "06 無骨雞排", "09 咕咾肉(玖)", "少油",
    ):
        assert removed not in text
    assert text.count("無異動") == 5


def test_daily_docx_uses_word_2007_compatible_cell_margin_xml():
    payload = render_daily_change_sheet_docx(_minimal_daily({
        "meal_label": "午餐", "has_changes": True,
        "summary": {
            "replacement_group_count": 1, "replacement_item_count": 1,
            "manual_acknowledgement_count": 0, "requires_reconfirmation_count": 0,
        },
        "replacement_groups": [{
            "replacement_dish": {"code": "R-CODE", "name": "替代菜"},
            "items": [{
                "current_room": "505", "case_name": "王小姐", "case_number": "CASE-CODE",
                "original_dish": {"code": "O-CODE", "name": "原菜"},
                "replacement_dish": {"code": "R-CODE", "name": "替代菜"},
                "restriction_groups": [{"name": "不牛"}], "status": "replaced", "warnings": [],
            }],
        }],
        "manual_acknowledgements": [], "requires_reconfirmation": [], "warnings": [],
    }))
    with ZipFile(BytesIO(payload)) as archive:
        assert archive.testzip() is None
        xml = archive.read("word/document.xml")
        ElementTree.fromstring(xml)
    assert b"<w:left" in xml and b"<w:right" in xml
    assert b"<w:start" not in xml and b"<w:end" not in xml
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for properties in root.iter(f"{namespace}tcPr"):
        names = [child.tag.removeprefix(namespace) for child in properties]
        if "shd" in names:
            assert names.index("shd") < names.index("tcMar")
        assert names.index("tcMar") < names.index("vAlign")


def test_daily_docx_title_and_weekday_share_one_twenty_point_header_row():
    meal = {
        "meal_label": "午餐", "has_changes": False,
        "summary": {
            "replacement_group_count": 0, "replacement_item_count": 0,
            "manual_acknowledgement_count": 0, "requires_reconfirmation_count": 0,
        },
        "replacement_groups": [], "manual_acknowledgements": [],
        "requires_reconfirmation": [], "warnings": [],
    }
    thursday = Document(BytesIO(render_daily_change_sheet_docx(_minimal_daily(meal))))
    header = thursday.tables[0]
    assert len(header.rows) == 1 and len(header.rows[0].cells) == 2
    title_cell, date_cell = header.rows[0].cells
    assert title_cell.text == "月子餐每日異動單"
    assert date_cell.text == "日期：2026/09/10（星期四）"
    title_run = next(run for run in title_cell.paragraphs[0].runs if run.text)
    date_run = next(run for run in date_cell.paragraphs[0].runs if run.text)
    assert title_run.bold is True and title_run.font.size.pt == 20
    assert date_run.bold is True and date_run.font.size.pt == 20
    assert title_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert date_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    borders = header._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert borders.find(qn("w:bottom")).get(qn("w:val")) == "single"
    for edge in ("top", "left", "right", "insideH", "insideV"):
        assert borders.find(qn(f"w:{edge}")).get(qn("w:val")) == "nil"

    monday_payload = _minimal_daily(meal)
    monday_payload["target_date"] = "2026-09-07"
    sunday_payload = _minimal_daily(meal)
    sunday_payload["target_date"] = "2026-09-13"
    monday_text = _all_text(Document(BytesIO(render_daily_change_sheet_docx(monday_payload))))
    sunday_text = _all_text(Document(BytesIO(render_daily_change_sheet_docx(sunday_payload))))
    assert "日期：2026/09/07（星期一）" in monday_text
    assert "日期：2026/09/13（星期日）" in sunday_text


def test_daily_docx_stale_replacement_marks_only_actual_stale_item(client, db_session):
    headers, _, _, first_items, _, _, acknowledgement = _configured_handlings(client, db_session)
    assert client.post(
        f"/api/v1/postpartum/conflict-acknowledgements/{acknowledgement['id']}/cancel",
        headers=headers,
    ).status_code == 204
    db_session.execute(delete(MenuDish).where(MenuDish.id == first_items[0]["original_menu_dish_id"]))
    db_session.commit()
    daily = client.get(f"/api/v1/postpartum/change-sheet/daily?target_date={TARGET_DATE}", headers=headers).json()
    lunch = next(item for item in daily["meals"] if item["postpartum_meal"] == "lunch")
    stale_reason = lunch["requires_reconfirmation"][0]["warnings"][0]["message"]
    document = _document(client.get(_path(), headers=headers))
    text = _all_text(document)
    assert "此替代內容尚需營養師重新確認" in text
    assert "⚠ 待重新確認" in text
    assert stale_reason in text
    details = next(table for table in document.tables if table.rows[0].cells[0].text == "床號" and table.rows[0].cells[3].text == "替代菜")
    assert sum("待重新確認" in row.cells[4].text for row in details.rows[1:]) == 1


def test_daily_docx_stale_acknowledgement_is_not_presented_as_complete(client, db_session):
    headers, _, _, _, second_items, group, _ = _configured_handlings(client, db_session)
    assert client.post(f"/api/v1/postpartum/replacement-groups/{group['id']}/cancel", headers=headers).status_code == 204
    db_session.execute(delete(MenuDish).where(MenuDish.id == second_items[0]["original_menu_dish_id"]))
    db_session.commit()
    daily = client.get(f"/api/v1/postpartum/change-sheet/daily?target_date={TARGET_DATE}", headers=headers).json()
    lunch = next(item for item in daily["meals"] if item["postpartum_meal"] == "lunch")
    stale_reason = lunch["requires_reconfirmation"][0]["warnings"][0]["message"]
    document = _document(client.get(_path(), headers=headers))
    acknowledgement = next(
        table for table in document.tables
        if len(table.rows[0].cells) > 3 and table.rows[0].cells[3].text == "結果"
    )
    assert acknowledgement.rows[1].cells[3].text == "⚠ 待重新確認"
    text = _all_text(document)
    assert "原處理：人工確認不需替代" in text
    assert stale_reason in text


def test_daily_docx_export_never_commits(client, db_session, monkeypatch):
    headers = auth(client, db_session)

    def reject_commit():
        raise AssertionError("DOCX read endpoint must not commit")

    monkeypatch.setattr(db_session, "commit", reject_commit)
    assert client.get(_path(), headers=headers).status_code == 200
